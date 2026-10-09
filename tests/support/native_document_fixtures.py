"""Small native originals created in memory; never production source files."""

from io import BytesIO

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def docx_business_qa():
    from docx import Document

    document = Document()
    document.add_heading("公司投资者关系活动记录", 0)
    document.add_paragraph("问：公司新产品海外客户认证有何进展？")
    document.add_paragraph(
        "答：公司新产品已通过海外客户认证并获得重复订单，进入量产交付。"
    )
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "业务"
    table.cell(0, 1).text = "最新进展"
    table.cell(1, 0).text = "新业务"
    table.cell(1, 1).text = "新产品获得客户重复订单，并扩大海外生产产能。"
    out = BytesIO()
    document.save(out)
    return out.getvalue()


def natural_html_call():
    return (
        "<html><head><title>Example Quarterly Earnings Conference Call</title></head>"
        "<body><nav>Transcript</nav><main><h2>Transcript</h2>"
        "<p>Example Quarterly Earnings Conference Call</p><p>JANE DOE:</p>"
        "<p>We launched a new product and expanded overseas capacity for customers.</p>"
        "<p>JANE DOE: We will now move over to Q&amp;A.</p><p>ALEX SMITH:</p>"
        "<p>Is customer validation complete for the new product?</p><p>JANE DOE:</p>"
        "<p>Customer validation remains planned next year, not completed.</p>"
        "<p>AMY ROE:</p><p>Thank you.</p><p>(Operator Direction.)</p><p>END</p>"
        "</main><footer>We launched a new unrelated retail site.</footer></body></html>"
    ).encode("utf-8")
