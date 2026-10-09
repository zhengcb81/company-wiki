"""Responsibility tests: ordinal identity, offline image evidence and replay."""

from dataclasses import replace
import hashlib
import io
import json
import time
import zipfile
from xml.etree import ElementTree as ET

import pytest
from company_wiki import document_normalization as dn
from conftest import build_pptx, sha256_hex, source_id_for, html_page

MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def normalize(data, **kwargs):
    return dn.normalize_document(
        data,
        source_id=source_id_for(data),
        source_sha256=sha256_hex(data),
        mime_type=MIME,
        **kwargs,
    )


def ordered_deck(reverse=False):
    data = build_pptx(with_group=True)
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as target:
        for name in source.namelist():
            payload = source.read(name)
            if name == "ppt/presentation.xml":
                xml = ET.fromstring(payload)
                rows = xml.find(
                    "{http://schemas.openxmlformats.org/presentationml/2006/main}sldIdLst"
                )
                rows[0].set("id", "256")
                rows[1].set("id", "260")
                if reverse:
                    first = rows[0]
                    rows.remove(first)
                    rows.append(first)
                payload = ET.tostring(xml, encoding="utf-8")
            target.writestr(name, payload)
    return out.getvalue()


class FakeOCR:
    def __init__(
        self, text="We launched a new product", confidence=0.95, blank=False, fail=None
    ):
        self.calls = 0
        self.text = text
        self.confidence = confidence
        self.blank = blank
        self.fail = fail
        self.identity_manifest = {"adapter": "fake-local-ocr/1", "text": text}
        self.fingerprint = hashlib.sha256(
            json.dumps(
                self.identity_manifest, sort_keys=True, separators=(",", ":")
            ).encode()
        ).hexdigest()

    def transcribe(self, data, *, limits):
        from company_wiki.document_normalization.local_ocr import (
            ImageOCRResult,
            OCRLine,
            LocalOCRError,
        )

        self.calls += 1
        if self.fail:
            raise LocalOCRError(self.fail)
        return ImageOCRResult(
            1,
            1,
            ()
            if self.blank
            else (OCRLine(self.text, (0.0, 0.0, 1.0, 1.0), self.confidence),),
        )


def test_current_pptx_uses_display_ordinal_and_separate_package_id():
    doc = normalize(ordered_deck())
    assert {u.coordinates.page_number for u in doc.units} == {1, 2}
    assert doc.parser_version == "1.1.0"
    assert {u.metadata["slide_id"] for u in doc.units} == {256, 260}
    assert all(
        u.metadata["source_locator"].startswith("cwp-pptx-shape/2|") for u in doc.units
    )
    reverse = normalize(ordered_deck(reverse=True))
    assert reverse.units[0].raw_text == "Grouped caption A"
    assert reverse.units[0].coordinates.page_number == 1
    assert reverse.units[0].metadata["slide_id"] == 260


def test_image_body_is_explicit_optional_and_deduplicated_with_separate_locations():
    data = build_pptx(with_picture_only_slide=True, with_duplicate_picture=True)
    assert normalize(data).structure.opaque_pages
    adapter = FakeOCR()
    doc = normalize(data, ocr=adapter)
    lines = [u for u in doc.units if u.unit_kind == "pptx_image_ocr_line"]
    assert adapter.calls == 1
    assert len(lines) == 3 and len({u.unit_id for u in lines}) == 3
    assert {u.coordinates.page_number for u in lines} == {2, 3}
    assert doc.parser_version == "2.0.0"
    assert all("ocr_used" in u.quality_flags for u in lines)
    assert all(u.metadata["coordinate_space"] == "image_pixels" for u in lines)
    assert doc.metadata["ocr"]["recall_status"] == "unverified"
    assert not doc.structure.coverage_complete


def test_legacy_pptx_replays_old_id_while_html_identity_is_unchanged():
    data = ordered_deck()
    legacy = normalize(data, parser_version="1.0.0")
    unit = legacy.units[0]
    assert unit.coordinates.page_number == 256
    assert unit.parser_version == "1.0.0"
    assert unit.metadata["source_locator"].startswith("cwp-pptx-shape/1|")
    assert (
        dn.replay_unit(
            data,
            source_sha256=sha256_hex(data),
            unit=unit,
            limits=dn.NormalizationLimits(),
        )
        == unit.raw_text
    )
    current_global_adapter = FakeOCR()
    assert (
        dn.replay_unit(
            data,
            source_sha256=sha256_hex(data),
            unit=unit,
            limits=dn.NormalizationLimits(),
            ocr=current_global_adapter,
        )
        == unit.raw_text
    )
    assert current_global_adapter.calls == 0
    html = html_page("<p>We launched a new product</p>").encode()
    doc = dn.normalize_document(
        html,
        source_id=source_id_for(html),
        source_sha256=sha256_hex(html),
        mime_type="text/html",
    )
    assert doc.parser_version == "1.0.0"
    assert doc.units[0].metadata["source_locator"].startswith("cwp-html-dom/1|")


@pytest.mark.parametrize("mode", ["blank", "low", "failure"])
def test_recognition_quality_never_becomes_complete_or_skip(mode):
    adapter = FakeOCR(
        blank=mode == "blank",
        confidence=0.6 if mode == "low" else 0.95,
        fail="OCR_IMAGE_INVALID" if mode == "failure" else None,
    )
    doc = normalize(build_pptx(with_picture_only_slide=True), ocr=adapter)
    page = next(p for p in doc.metadata["ocr"]["pages"] if p["page_number"] == 2)
    assert not doc.structure.coverage_complete
    assert (
        page["status"]
        == {"blank": "empty", "low": "low_confidence", "failure": "failed"}[mode]
    )
    if mode == "low":
        assert "low_ocr_confidence" in doc.units[-1].quality_flags


def test_multi_span_replay_uses_one_normalization_and_rejects_tampered_metadata():
    data = build_pptx(with_picture_only_slide=True, with_duplicate_picture=True)
    doc = normalize(data, ocr=FakeOCR())
    selected = [u for u in doc.units if u.unit_kind == "pptx_image_ocr_line"]
    adapter = FakeOCR()
    assert dn.replay_units(
        data,
        source_sha256=sha256_hex(data),
        units=selected,
        limits=dn.NormalizationLimits(),
        ocr=adapter,
    ) == tuple(u.raw_text for u in selected)
    assert adapter.calls == 1
    span = selected[0].to_evidence_span(
        topics=["business_progress"], selection_reasons=["body"]
    )
    assert (
        dn.replay_evidence_spans(
            data,
            source_id=source_id_for(data),
            source_sha256=sha256_hex(data),
            mime_type=MIME,
            evidence_spans=(span,),
            limits=dn.NormalizationLimits(),
            ocr=FakeOCR(),
        )
        == 1
    )
    for key, value in (
        ("media_sha256", "f" * 64),
        ("ocr_fingerprint", "e" * 64),
        ("ocr_confidence", 0.8),
    ):
        bad = replace(selected[0], metadata={**selected[0].metadata, key: value})
        with pytest.raises(ValueError):
            dn.replay_units(
                data,
                source_sha256=sha256_hex(data),
                units=(bad,),
                limits=dn.NormalizationLimits(),
                ocr=FakeOCR(),
            )


def test_deadline_pixel_text_and_unit_limits_are_named_and_do_not_publish():
    from company_wiki.document_normalization.local_ocr import OCRLimits

    data = build_pptx(with_picture_only_slide=True)
    with pytest.raises(dn.NormalizationLimitError):
        normalize(
            data,
            ocr=FakeOCR(),
            limits=dn.NormalizationLimits(deadline=time.monotonic() - 1),
        )
    with pytest.raises(dn.NormalizationLimitError):
        normalize(data, ocr=FakeOCR(), limits=dn.NormalizationLimits(max_units=2))
    with pytest.raises(dn.NormalizationLimitError):
        normalize(
            data, ocr=FakeOCR(), limits=dn.NormalizationLimits(max_text_output_bytes=70)
        )
    # two occurrences count image processing bounds, unique media is inferred once
    with pytest.raises(dn.NormalizationLimitError):
        normalize(
            build_pptx(with_picture_only_slide=True, with_duplicate_picture=True),
            ocr=FakeOCR(),
            ocr_limits=OCRLimits(max_images=1),
        )


def test_ten_span_replay_infers_only_two_selected_verified_media_and_rejects_forgery():
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches
    from company_wiki.document_normalization.local_ocr import ImageOCRResult, OCRLine

    presentation = Presentation()
    for color in ("red", "green", "blue"):
        png = io.BytesIO()
        Image.new("RGB", (1, 1), color).save(png, format="PNG")
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(io.BytesIO(png.getvalue()), Inches(0), Inches(0))
    output = io.BytesIO()
    presentation.save(output)
    data = output.getvalue()

    class FiveLines(FakeOCR):
        def transcribe(self, data, *, limits):
            self.calls += 1
            return ImageOCRResult(
                1,
                1,
                tuple(
                    OCRLine("Body line " + str(i), (0, 0, 1, 1), 0.95) for i in range(5)
                ),
            )

    initial = FiveLines()
    doc = normalize(data, ocr=initial)
    assert initial.calls == 3
    selected = tuple(u for u in doc.units if u.coordinates.page_number in {1, 2})
    spans = tuple(
        u.to_evidence_span(
            topics=["business_progress"], selection_reasons=["selected_media"]
        )
        for u in selected
    )
    assert len(spans) == 10
    adapter = FiveLines()
    assert (
        dn.replay_evidence_spans(
            data,
            source_id=source_id_for(data),
            source_sha256=sha256_hex(data),
            mime_type=MIME,
            evidence_spans=spans,
            limits=dn.NormalizationLimits(),
            ocr=adapter,
        )
        == 10
    )
    assert adapter.calls == 2
    good = spans[0]
    locator = good.structured_value["source_locator"]
    for changed in (
        locator.replace("|p=0|", "|p=1|"),
        locator.replace(good.structured_value["media_sha256"], "f" * 64),
        locator.replace("|b=0.000,0.000,1.000,1.000", "|b=0.000,0.000,0.500,1.000"),
    ):
        bad = good.__class__.create(
            source_id=good.source_id,
            coordinates=good.coordinates,
            raw_text=good.raw_text,
            structured_value={**good.structured_value, "source_locator": changed},
            parser_name=good.parser_name,
            parser_version=good.parser_version,
            parse_status=good.parse_status,
            quality_flags=good.quality_flags,
        )
        with pytest.raises(ValueError):
            dn.replay_evidence_spans(
                data,
                source_id=source_id_for(data),
                source_sha256=sha256_hex(data),
                mime_type=MIME,
                evidence_spans=(bad,),
                limits=dn.NormalizationLimits(),
                ocr=FiveLines(),
            )


@pytest.mark.parametrize("version", ["", "9.9.9"])
def test_explicit_unknown_pptx_version_does_not_silently_choose_current(version):
    with pytest.raises(ValueError, match="parser_version"):
        normalize(build_pptx(), parser_version=version)
