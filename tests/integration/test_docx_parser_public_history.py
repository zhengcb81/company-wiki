"""Real DOCX parser upgrade through finite Worker, public read and cache.

The old producer bootstrap selects its runtime parser before processing; saved
jobs, bundles, manifests, locators and budget rows are never edited.
"""
from contextlib import closing
from io import BytesIO
import json
import sqlite3
import subprocess
import sys

from docx import Document

from company_wiki.automation.narrative_batch import _thaw_generations
from company_wiki.automation.narrative_contracts import NarrativeBundle
from company_wiki.automation.narrative_run_store import NarrativeRunStore
from company_wiki.document_normalization.units import DOCX_PARSER_VERSION
from support import narrative_batch_fixtures as raw_fixtures
from support import official_json_batch_fixture as fixtures


official_json_loopback_model = fixtures.official_json_loopback_model
CONTEXT = "The return calculation has not changed"
QUALIFIER = "Efficiency depends on workload mix, token usage and silicon price performance."
BUSINESS = "Company launched a new product and expanded overseas production capacity."


def _original():
    document = Document()
    document.add_heading(CONTEXT, level=1)
    document.add_heading(QUALIFIER, level=1)
    document.add_paragraph(BUSINESS)
    out = BytesIO()
    document.save(out)
    return out.getvalue()


def _dump(path):
    with closing(sqlite3.connect("file:" + path.as_posix() + "?mode=ro", uri=True)) as db:
        return tuple(db.iterdump())


def _old_producer(state, path, run_id):
    # Only the parent selects the old version. Actual spawned Workers receive
    # that frozen parser identity and execute the supported old implementation.
    bootstrap = "\n".join((
        "from company_wiki.source_catalog.narrative_normalization import NarrativeNormalization as N",
        "old_identity = N.identity",
        "def identity(self, mime_type, *, document_id=None, parser_version=None):",
        "    if mime_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document' and parser_version is None:",
        "        parser_version = '1.0.0'",
        "    return old_identity(self, mime_type, document_id=document_id, parser_version=parser_version)",
        "N.identity = identity",
        "from company_wiki.automation.narrative_batch_cli import main",
        "raise SystemExit(main())",
    ))
    command = [sys.executable, "-B", "-c", bootstrap,
               "--project-root", str(state.root), "--catalog-config", str(state.config_path),
               "--automation-db", str(state.automation_db),
               "--work-dir", str(state.root / ("work-" + run_id)), "--request", str(path)]
    process = subprocess.run(command, cwd=state.root, env=fixtures.subprocess_env(),
                             capture_output=True, text=True, encoding="utf-8", timeout=60)
    assert fixtures.KEY not in process.stdout + process.stderr
    result = json.loads(process.stdout)
    state.calls.append({"command": command, "exit_code": process.returncode,
                        "receipt": result, "stderr": process.stderr})
    return process, result


def _read(state, receipt, version, *, context_selected):
    assert receipt["status"] == "completed" and len(receipt["items"]) == 1, receipt
    item = receipt["items"][0]
    assert item["status"] == "completed" and item["errors"] == [], item
    reference = item["artifact_ref"]
    wire, replay = fixtures.invoke_read(state, reference)
    bundle = NarrativeBundle.from_dict(wire)
    assert bundle.versions.parser == version
    assert all(span.parser_name == "cwp_document_normalization" for span in bundle.evidence_spans)
    assert replay["replay_status"] == "verified" and bundle.summary.translate is False
    assert bundle.source_ref.to_dict() == state.raw_ref
    texts = {span.raw_text for span in bundle.evidence_spans}
    assert BUSINESS in texts and QUALIFIER in texts
    assert (CONTEXT in texts) is context_selected
    assert all(span.parser_version == version for span in bundle.evidence_spans)
    if not context_selected:
        qualifier = next(span for span in bundle.evidence_spans if span.raw_text == QUALIFIER)
        assert qualifier.structured_value["unit_kind"] == "docx_heading"
        assert not any(span.structured_value.get("selection_group_id")
                       for span in bundle.evidence_spans)
    run = NarrativeRunStore(state.automation_db).get_run(receipt["run_id"])
    binding = json.loads(run.binding_json)
    manifests = _thaw_generations(binding)
    assert len(manifests) == 1
    assert next(iter(manifests.values()))["parser_component"]["version"] == version
    return reference, binding


def test_docx_new_parser_preserves_old_public_history_and_creates_distinct_generation(
    tmp_path, monkeypatch, official_json_loopback_model,
):
    server = official_json_loopback_model
    data = _original()
    monkeypatch.setattr(raw_fixtures, "source_documents", lambda **kwargs: [
        ("en", "business-ir.docx", "Company business communication", "investor_relations", data),
    ])
    with fixtures.official_batch_state(tmp_path) as state:
        items = [{"kind": "raw", "source_ref": state.raw_ref}]
        old_path = fixtures.request_for(state, server.endpoint, "docx-old", items=items)
        process, old = _old_producer(state, old_path, "docx-old")
        assert process.returncode == 0, (old, process.stderr)
        old_ref, old_binding = _read(state, old, "1.0.0", context_selected=True)
        assert len(server.requests) == 1 and server.errors == []
        before_auto = _dump(state.automation_db)
        before_source = _dump(state.catalog.config.database_path)
        process, resumed = fixtures.invoke_batch(state, old_path, "docx-old")
        assert process.returncode == 0, (resumed, process.stderr)
        assert resumed["items"] == old["items"] and resumed["budget"] == old["budget"]
        assert _read(state, resumed, "1.0.0", context_selected=True) == (old_ref, old_binding)
        assert _dump(state.automation_db) == before_auto
        assert _dump(state.catalog.config.database_path) == before_source
        assert len(server.requests) == 1
        current_path = fixtures.request_for(state, server.endpoint, "docx-current", items=items)
        process, current = fixtures.invoke_batch(state, current_path, "docx-current")
        assert process.returncode == 0, (current, process.stderr)
        new_ref, _ = _read(state, current, DOCX_PARSER_VERSION, context_selected=False)
        assert DOCX_PARSER_VERSION == "1.1.0"
        assert new_ref["artifact_version_id"] != old_ref["artifact_version_id"]
        assert len(server.requests) == 2 and server.errors == []
        assert old["budget"]["tokens"] == current["budget"]["tokens"] == 92
        reuse_path = fixtures.request_for(state, server.endpoint, "docx-reuse", items=items)
        process, reused = fixtures.invoke_batch(state, reuse_path, "docx-reuse")
        assert process.returncode == 0
        reused_ref, _ = _read(state, reused, DOCX_PARSER_VERSION, context_selected=False)
        assert reused_ref == new_ref
        assert reused["budget"]["tokens"] == reused["budget"]["estimated_micro_usd"] == 0
        assert len(server.requests) == 2 and server.errors == []
        print("DOCX-PUBLIC-HISTORY " + json.dumps({
            "old_parser": "1.0.0", "new_parser": DOCX_PARSER_VERSION,
            "old_reference": old_ref, "new_reference": new_ref,
            "old_budget": old["budget"], "new_budget": current["budget"],
            "reuse_budget": reused["budget"], "loopback_posts": len(server.requests),
            "actual_cli_calls": [{"exit_code": call["exit_code"], "run_id": call["receipt"]["run_id"]}
                                 for call in state.calls],
            "original_sha256": fixtures.sha(data),
            "config_sha256": fixtures.sha(state.config_bytes),
        }, sort_keys=True))
