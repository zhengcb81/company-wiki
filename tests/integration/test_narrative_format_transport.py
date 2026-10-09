"""Actual catalog/AUTO/handler/projector/read integration, in disposable lakes."""
from __future__ import annotations

import hashlib
from io import BytesIO
import json
import os
from pathlib import Path

import pytest

from company_wiki.automation.narrative_transport import NarrativeTransportReader
from company_wiki.automation.narrative_transport_contracts import NarrativeReadRequest
from support.narrative_transport_fixture import published_fixture
from company_wiki.document_normalization.units import (
    PARSER_VERSION as HTML_PARSER_VERSION,
    PPTX_PARSER_VERSION,
)


BUSINESS = "Company launched a new product and expanded overseas capacity for new customers."


def _transport(tmp_path, data, suffix):
    before = tuple(tmp_path.iterdir())
    with published_fixture(tmp_path, source_spec={
        "name": "source." + suffix, "data": data, "language": "en",
        "title": "Company annual business update", "document_kind": "annual_report",
    }) as fixture:
        reader = NarrativeTransportReader(fixture.artifacts, fixture.reader)
        reference = reader.reference(fixture.source_ref)
        request = fixture.read_request(reference)
        result = reader.read(NarrativeReadRequest.from_dict(request))
        assert result.data == fixture.payload
        payload = json.loads(result.data)
        assert payload["evidence_spans"]
        assert payload["source_metadata"]["language"] == "en"
        assert payload["summary"]["translate"] is False
        assert result.receipt["replay_status"] == "verified"
        assert result.receipt["locator_count"] == len(payload["evidence_spans"])
        expected_parser = PPTX_PARSER_VERSION if suffix == "pptx" else HTML_PARSER_VERSION
        assert payload["versions"]["parser"] == expected_parser
    assert tuple(tmp_path.iterdir()) == before


@pytest.mark.parametrize("suffix", ["htm", "html", "pptx"])
def test_format_real_auto_publication_and_public_transport_restores_test_lake(tmp_path, suffix):
    if suffix == "pptx":
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(7), Inches(2)).text = BUSINESS
        buffer = BytesIO()
        deck.save(buffer)
        data = buffer.getvalue()
    else:
        data = f"<html><body><p>{BUSINESS}</p></body></html>".encode()
    _transport(tmp_path, data, suffix)


def test_real_sec_html_uses_the_same_pipeline_and_preserves_original(tmp_path):
    sample = os.environ.get("CWP_R6_HTML_SAMPLE")
    if not sample:
        pytest.skip("explicit read-only real original not provided")
    path = Path(sample)
    data = path.read_bytes()
    expected = "99d693f6c1544144ebeee92954f151a85bc62111837530a42855953bc01d0bbe"
    assert hashlib.sha256(data).hexdigest() == expected
    _transport(tmp_path, data, "htm")
    assert hashlib.sha256(path.read_bytes()).hexdigest() == expected
