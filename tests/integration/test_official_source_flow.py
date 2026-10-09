"""W03 responsibilities: immutable official import and honest bounded discovery."""

from contextlib import contextmanager
import hashlib
import json
from pathlib import Path
import tempfile
import socket
from support import narrative_batch_fixtures as batch_fixtures

# Capture before the hermetic autouse fixture. Windows asyncio uses a local
# socket pair even for an httpx MockTransport; only that local pair is allowed.
import pytest
from company_wiki.source_catalog.models import CatalogConfig, RootSpec
from company_wiki.source_catalog.service import SourceCatalog
from company_wiki.source_catalog.source_reader import SourceVersionReader


_SOCKET_CONNECT = socket.socket.connect


@contextmanager
def owned_lake():
    with tempfile.TemporaryDirectory(prefix="cwpaW03-") as folder:
        root = Path(folder)
        (root / "companies").mkdir()
        catalog = SourceCatalog(
            CatalogConfig(
                project_root=root,
                catalog_dir=root / "catalog",
                roots=(RootSpec("company_raw", root / "companies", "company_raw"),),
            )
        )
        try:
            yield root, catalog
        finally:
            catalog.close()
    assert not root.exists()


def request(
    data,
    *,
    entity="Acme",
    kind="investor_relations",
    mime="text/html",
    published="2026-08-26",
    language="en",
):
    return {
        "schema_version": "official-source-import-request/1",
        "request_id": "w03-test-1",
        "source": {
            "entity": entity,
            "market": "US",
            "security_id": "ACME" if entity == "Acme" else None,
            "document_kind": kind,
            "title": entity + " business update",
            "publisher": entity,
            "source_url": "https://official.example/" + entity + "/update",
            "published_date": published,
            "fiscal_year": None,
            "fiscal_period": None,
            "language": language,
        },
        "content_sha256": hashlib.sha256(data).hexdigest(),
        "mime_type": mime,
        "max_bytes": 1048576,
        "capture_receipt": {
            "capture_method": "local_document",
            "tool_name": "w03-fixture",
            "tool_call_id": "fixture-1",
            "captured_at": "2026-10-09T00:00:00Z",
            "response_bytes": len(data),
            "content_sha256": hashlib.sha256(data).hexdigest(),
        },
    }


def _import(catalog, data, **kw):
    from company_wiki.source_catalog.official_source_flow import import_official_source

    return import_official_source(catalog, original=data, request=request(data, **kw))


def test_official_import_is_pathless_byte_verified_and_idempotent():
    data = b"<html><h1>Acme business update</h1><p>We launched a new product and expanded overseas capacity.</p></html>"
    with owned_lake() as (root, catalog):
        first = _import(catalog, data)
        assert first["schema_version"] == "official-source-import-result/1"
        ref = SourceVersionReader(catalog).query_ref(
            **{k: first["source_ref"][k] for k in ["document_id", "source_id"]},
            expected_sha256=first["source_ref"]["content_sha256"],
        )
        assert (
            SourceVersionReader(catalog).open_version(ref, purpose="source_export").data
            == data
        )
        before = {
            p: p.read_bytes() for p in (root / "companies").rglob("*") if p.is_file()
        }
        other = request(data)
        other["request_id"] = "another-attempt"
        other["capture_receipt"]["tool_call_id"] = "another-capture"
        from company_wiki.source_catalog.official_source_flow import (
            import_official_source,
        )

        again = import_official_source(catalog, original=data, request=other)
        assert (
            again["status"] == "deduplicated"
            and first["source_ref"] == again["source_ref"]
        )
        assert before == {
            p: p.read_bytes() for p in (root / "companies").rglob("*") if p.is_file()
        }
        assert "canonical_path" not in json.dumps(first) and str(
            root
        ) not in json.dumps(first)
        assert len([p for p in (root / "companies").rglob("*.html")]) == 1
        changed = _import(catalog, data.replace(b"new product", b"second new product"))
        assert changed["source_ref"] != first["source_ref"]
        assert all(p.read_bytes() == b for p, b in before.items())


@pytest.mark.parametrize("entity", ["Microsoft", "OpenAI"])
def test_other_company_and_unknown_dates_are_never_relabelled(entity):
    with owned_lake() as (_, catalog):
        result = _import(
            catalog,
            f"<html><h1>{entity}</h1><p>Customer demand exceeds available supply.</p></html>".encode(),
            entity=entity,
            published=None,
            language="unknown",
        )
        m = result["metadata"]
        assert (
            m["entity"] == entity
            and m["published_date"] is None
            and m["fiscal_year"] is None
        )
        ref = result["source_ref"]
        reader = SourceVersionReader(catalog)
        actual = reader.describe_version(
            reader.query_ref(
                ref["document_id"], ref["source_id"], ref["content_sha256"]
            )
        )
        assert (
            actual["display_name"] == entity
            and actual["published_date"] is None
            and actual["fiscal_year"] is None
        )
        assert result["capture_receipt"]["captured_at"] != "" + str(
            actual["published_date"]
        )


@pytest.mark.parametrize(
    "data,mime,error",
    [
        (
            b"<html><title>Access Denied</title><body>Access Denied</body></html>",
            "text/html",
            "response_is_error",
        ),
        (b"%PDF-1.4\n%%EOF", "application/pdf", "invalid_pdf"),
        (b"<html>not a PDF</html>", "application/pdf", "mime_mismatch"),
    ],
)
def test_bad_originals_refused_before_canonical_write(data, mime, error):
    from company_wiki.source_catalog.official_source_flow import OfficialSourceError

    with owned_lake() as (root, catalog):
        with pytest.raises(OfficialSourceError, match=error):
            _import(catalog, data, mime=mime)
        assert not list((root / "companies").rglob("*"))
        assert not list((root / "catalog/staging").glob("*"))


def test_bad_sha_refused_and_failed_import_cleans_only_owned_staging(monkeypatch):
    from company_wiki.source_catalog.official_source_flow import (
        import_official_source,
        OfficialSourceError,
    )

    data = b"CEO: We launched a new product and expanded overseas capacity."
    with owned_lake() as (root, catalog):
        req = request(data, mime="text/plain", kind="investor_call_transcript")
        req["content_sha256"] = "0" * 64
        with pytest.raises(OfficialSourceError, match="source_sha_mismatch"):
            import_official_source(catalog, original=data, request=req)
        assert not list((root / "companies").rglob("*"))


def test_discovery_notice_dynamic_future_and_unknown_are_explicit():
    from company_wiki.source_catalog.official_discovery import (
        discover_official_documents,
    )

    html = """<html><h1>Acme investor relations</h1><div>Loading</div>
    <article><time datetime="2026-09-10"></time><a href="/notice">Upcoming conference event notice</a></article>
    <article><time datetime="2026-10-21"></time><a href="/speech.pdf">Investor presentation business expansion</a></article>
    <article><time datetime="2026-08-26"></time><a href="/call.pdf">Q2 earnings call transcript</a></article>
    <a href="/unknown.pdf">Business presentation</a></html>""".encode()
    out = discover_official_documents(
        [
            {
                "url": "https://official.example/ir",
                "original": html,
                "mime_type": "text/html",
                "capture_id": "real-fixture-receipt",
            }
        ],
        entity="Acme",
        as_of_date="2026-10-09",
        start_date="2026-02-01",
        max_items=30,
        max_bytes=8388608,
    )
    by = {x["source_url"]: x for x in out["candidates"]}
    assert by["https://official.example/notice"]["disposition"] == "index_notice"
    assert (
        by["https://official.example/call.pdf"]["document_kind"]
        == "investor_call_transcript"
    )
    assert (
        by["https://official.example/unknown.pdf"]["published_date"] is None
        and by["https://official.example/unknown.pdf"]["disposition"] == "date_unknown"
    )
    assert "https://official.example/speech.pdf" not in by
    assert (
        out["coverage_status"] == "partial"
        and "dynamic_content_unresolved" in out["limitations"]
    )
    assert out["receipt"]["event_ids"] == ["real-fixture-receipt"]


def test_discovery_limits_do_not_turn_truncation_into_no_materials():
    from company_wiki.source_catalog.official_discovery import (
        discover_official_documents,
    )

    body = b'<html><a href="/one.pdf">Earnings call</a><a href="/two.pdf">Investor business presentation</a></html>'
    out = discover_official_documents(
        [
            {
                "url": "https://official.example/ir",
                "original": body,
                "mime_type": "text/html",
                "capture_id": "fixture",
            }
        ],
        entity="Acme",
        as_of_date="2026-10-09",
        start_date="2026-02-01",
        max_items=1,
        max_bytes=8388608,
    )
    assert (
        out["coverage_status"] == "partial"
        and "item_limit_reached" in out["limitations"]
        and len(out["candidates"]) == 1
    )


def test_pdf_call_is_normalized_as_pdf_and_unknown_language_is_detected():
    from company_wiki.automation.narrative_batch import build_batch_events
    from company_wiki.automation.narrative_batch_request import NarrativeBatchRequest
    import fitz

    doc = fitz.open()
    doc.new_page().insert_text(
        (72, 72),
        "CEO: Company launched a new product and expanded overseas capacity for customers.",
    )
    data = doc.tobytes()
    doc.close()
    with owned_lake() as (_, catalog):
        result = _import(
            catalog,
            data,
            mime="application/pdf",
            kind="investor_call_transcript",
            language="unknown",
        )
        req = NarrativeBatchRequest.from_dict(
            {
                "schema_version": "narrative-batch-request/1",
                "run_id": "pdf-official",
                "sources": [result["source_ref"]],
                "profile": "P1",
                "max_seconds": 10,
                "max_tokens": 10000,
                "max_cost_usd": "1",
                "model": {
                    "model_id": "fixture",
                    "endpoint": "http://127.0.0.1:1/v1/chat/completions",
                    "api_key_env": "FAKE",
                    "allow_local_http": True,
                },
                "pricing": {
                    "version": "test/1",
                    "input_micro_usd_per_million_tokens": 0,
                    "output_micro_usd_per_million_tokens": 0,
                },
            }
        )
        binding = build_batch_events(
            req, SourceVersionReader(catalog), now="2026-10-09T00:00:00Z"
        )
        payload = json.loads(binding.events[0].payload_json)
        assert (
            payload["source_metadata"]["source_class"] == "filing"
        )  # carrier parser class; actual kind remains transcript
        assert payload["source_metadata"]["document_kind"] == "investor_call_transcript"
        assert payload["source_metadata"]["language"] == "en"


def test_format_selection_replay_32_spans_parses_original_once(monkeypatch):
    from company_wiki.document_normalization import normalize_document
    from company_wiki.source_catalog.narrative_evidence import select_narrative_evidence
    import company_wiki.automation.narrative_replay as replay
    from types import SimpleNamespace
    from company_wiki.source_contract import source_id_for_sha256
    from company_wiki.automation.narrative_formats import parser_component

    data = (
        "<html>"
        + "".join(
            "<p>Company launched a new product line "
            + str(n)
            + " and expanded overseas capacity for customers.</p>"
            for n in range(32)
        )
        + "</html>"
    ).encode()
    sha = hashlib.sha256(data).hexdigest()
    sid = source_id_for_sha256(sha)
    doc = normalize_document(
        data, source_id=sid, source_sha256=sha, mime_type="text/html"
    )
    package = select_narrative_evidence(
        doc.structure,
        title="Company business update",
        existing_kind="investor_relations",
    )
    assert len(package.evidence_spans) == 32
    calls = []
    original = replay.normalize_document

    def counted(*args, **kwargs):
        calls.append(1)
        return original(*args, **kwargs)

    monkeypatch.setattr(replay, "normalize_document", counted)
    name, version = parser_component("text/html")
    selected = SimpleNamespace(
        source_ref=SimpleNamespace(
            source_id=sid,
            content_sha256=sha,
            byte_size=len(data),
            mime_type="text/html",
        ),
        source_metadata=SimpleNamespace(source_class="filing"),
        versions=SimpleNamespace(parser=version),
        evidence_spans=package.evidence_spans,
    )
    assert replay.replay_narrative_evidence(data, selected) == 32 and len(calls) == 1
    selected.versions.parser = "different-parser-generation"
    with pytest.raises(replay.NarrativeReplayError):
        replay.replay_narrative_evidence(data, selected)


@pytest.mark.parametrize(
    "kind,mime,body",
    [
        (
            "investor_call_transcript",
            "text/plain",
            b"Full Conference Call Transcript\nCEO: We expanded overseas production capacity for customers.\n",
        ),
        (
            "investor_call_transcript",
            "application/json",
            b'[{"content":"Full Conference Call Transcript\\nCEO: We launched a new product and expanded overseas capacity."}]',
        ),
    ],
)
def test_txt_and_json_originals_stay_exact_and_real_kind(kind, mime, body):
    with owned_lake() as (_, catalog):
        result = _import(catalog, body, kind=kind, mime=mime)
        ref = result["source_ref"]
        reader = SourceVersionReader(catalog)
        assert (
            reader.open_version(
                reader.query_ref(
                    ref["document_id"], ref["source_id"], ref["content_sha256"]
                ),
                purpose="source_export",
            ).data
            == body
        )
        assert (
            reader.describe_version(
                reader.query_ref(
                    ref["document_id"], ref["source_id"], ref["content_sha256"]
                )
            )["document_kind"]
            == kind
        )


def test_public_import_cli_keeps_source_storage_path_private():
    import os
    import subprocess
    import sys

    repo = Path(__file__).resolve().parents[2]
    body = b"<html><h1>Acme</h1><p>Company launched a new product and expanded overseas capacity.</p></html>"
    with owned_lake() as (root, catalog):
        (root / "config").mkdir()
        cfg = root / "config/catalog.json"
        cfg.write_text(
            json.dumps(
                {
                    "schema_version": "1.0",
                    "catalog_dir": "catalog",
                    "roots": [
                        {
                            "root_id": "company_raw",
                            "path": "companies",
                            "kind": "company_raw",
                        }
                    ],
                }
            ),
            encoding="utf8",
        )
        raw = root / "incoming.html"
        raw.write_bytes(body)
        req = root / "request.json"
        req.write_text(json.dumps(request(body)), encoding="utf8")
        call = subprocess.run(
            [
                sys.executable,
                "-X",
                "utf8",
                "-B",
                "-m",
                "company_wiki.source_catalog.official_source_cli",
                "--config",
                str(cfg),
                "--input-file",
                str(raw),
                "--request",
                str(req),
            ],
            cwd=root,
            env={
                **os.environ,
                "PYTHONPATH": os.pathsep.join(
                    [str(repo / "src"), str(repo / "scripts")]
                ),
            },
            capture_output=True,
            text=True,
            encoding="utf8",
            timeout=20,
        )
        assert call.returncode == 0, call.stderr
        out = json.loads(call.stdout)
        assert (
            out["source_ref"]["schema_version"] == "2.0"
            and str(root) not in call.stdout
        )
        assert raw.read_bytes() == body


def test_official_index_http_uses_existing_shared_budget_with_mock_transport(
    monkeypatch,
):
    def local_pair_only(sock, address):
        if address[0] not in {"127.0.0.1", "::1"}:
            raise RuntimeError("external network forbidden in W03")
        return _SOCKET_CONNECT(sock, address)

    monkeypatch.setattr(socket.socket, "connect", local_pair_only)
    import httpx
    from company_wiki.source_catalog.official_discovery import fetch_official_indexes

    class Body(httpx.AsyncByteStream):
        async def __aiter__(self):
            yield b'<html><a href="/call.pdf">Q2 earnings call transcript</a></html>'

    async def respond(req):
        return httpx.Response(
            200, headers={"content-type": "text/html"}, stream=Body(), request=req
        )

    pages, receipts = fetch_official_indexes(
        ["https://official.example/ir"],
        transport=httpx.MockTransport(respond),
        max_bytes=200,
        max_seconds=5,
    )
    assert (
        len(pages) == 1
        and receipts[0]["status"] == "opened"
        and receipts[0]["usage_complete"]
    )
    assert receipts[0]["cumulative_response_bytes"] == len(pages[0]["original"])
    pages, receipts = fetch_official_indexes(
        ["https://official.example/ir"],
        transport=httpx.MockTransport(respond),
        max_bytes=10,
        max_seconds=5,
    )
    assert pages == [] and receipts[0]["status"] == "failed"


loopback_model_server = batch_fixtures.loopback_model_server


def _public_official_import(root, config, body, *, mime, kind, language):
    import os
    import subprocess
    import sys

    repo = Path(__file__).resolve().parents[2]
    identity = hashlib.sha256(body).hexdigest()
    original = root / ("incoming-" + identity[:16])
    original.write_bytes(body)
    value = request(body, mime=mime, kind=kind, language=language)
    request_path = root / ("import-" + identity[:16] + ".json")
    request_path.write_text(json.dumps(value), encoding="utf-8")
    call = subprocess.run(
        [
            sys.executable,
            "-X",
            "utf8",
            "-B",
            "-m",
            "company_wiki.source_catalog.official_source_cli",
            "--config",
            str(config),
            "--project-root",
            str(root),
            "--input-file",
            str(original),
            "--request",
            str(request_path),
        ],
        cwd=root,
        env={
            **os.environ,
            "PYTHONPATH": str(repo / "src"),
            "PYTHONDONTWRITEBYTECODE": "1",
        },
        capture_output=True,
        timeout=20,
    )
    assert call.returncode == 0, call.stderr
    assert str(root).encode() not in call.stdout
    assert original.read_bytes() == body
    return json.loads(call.stdout)


def _configured_invoke(root, config, req_path, llm_config):
    import os
    import subprocess
    import sys

    repo = Path(__file__).resolve().parents[2]
    env = {
        **os.environ,
        "PYTHONPATH": os.pathsep.join([str(repo / "src"), str(repo / "scripts")]),
        "PYTHONDONTWRITEBYTECODE": "1",
        "PYTHON_DOTENV_DISABLED": "1",
    }
    call = subprocess.run(
        [
            sys.executable,
            "-X",
            "utf8",
            "-B",
            str(repo / "scripts/narrative_batch_configured.py"),
            "--llm-config",
            str(llm_config),
            "--allow-local-model-http",
            "--project-root",
            str(root),
            "--catalog-config",
            str(config),
            "--automation-db",
            str(root / "AUTO.sqlite"),
            "--work-dir",
            str(root / "work"),
            "--request",
            str(req_path),
        ],
        cwd=root,
        env=env,
        capture_output=True,
        text=True,
        encoding="utf8",
        timeout=60,
    )
    assert batch_fixtures.KEY not in call.stdout + call.stderr
    assert "LEGACY WRITER BLOCKED" not in call.stdout + call.stderr
    assert call.stdout.strip(), call.stderr
    return call, json.loads(call.stdout)


def _public_narrative(root, config, operation, request_value):
    import os
    import subprocess
    import sys

    repo = Path(__file__).resolve().parents[2]
    call = subprocess.run(
        [
            sys.executable,
            "-X",
            "utf8",
            "-B",
            "-m",
            "company_wiki.source_catalog.narrative_transport_cli",
            "--config",
            str(config),
            "--operation",
            operation,
        ],
        cwd=root,
        env={
            **os.environ,
            "PYTHONPATH": str(repo / "src"),
            "PYTHONDONTWRITEBYTECODE": "1",
        },
        input=json.dumps(request_value).encode("utf-8"),
        capture_output=True,
        timeout=20,
    )
    assert call.returncode == 0, call.stderr
    assert batch_fixtures.KEY.encode() not in call.stdout + call.stderr
    receipt = json.loads(call.stderr)
    if operation == "read":
        reference = request_value["narrative_ref"]
        assert len(call.stdout) == reference["byte_size"]
        assert hashlib.sha256(call.stdout).hexdigest() == reference["artifact_sha256"]
        assert receipt["narrative_ref"] == reference
        proof_file = os.environ.get("CWPA_W03_PROOF_FILE")
        if proof_file and reference["source_ref"]["mime_type"] == "text/html":
            destination = Path(proof_file).parent
            (destination / "html-bundle.json").write_bytes(call.stdout)
            (destination / "html-read-receipt.json").write_bytes(call.stderr)
            (destination / "html-read-request.json").write_text(
                json.dumps(request_value, ensure_ascii=False, indent=2) + "\n",
                encoding="utf-8",
            )
    return json.loads(call.stdout), receipt


def test_configured_official_three_formats_summary_public_read_and_repeat(
    loopback_model_server, monkeypatch
):
    import fitz
    from company_wiki.source_catalog.official_source_flow import (
        prepare_official_narrative_request,
    )
    from company_wiki.automation.narrative_run_store import NarrativeRunStore

    monkeypatch.setenv("DEEPSEEK_API_KEY", batch_fixtures.KEY)
    doc = fitz.open()
    doc.new_page().insert_text(
        (72, 72),
        "CEO: We launched a new product and expanded overseas capacity for customers.",
    )
    pdf = doc.tobytes()
    doc.close()
    sources = [
        (
            b"<html><h1>Acme business update</h1><p>We launched a new product and expanded overseas capacity for customers.</p></html>",
            "text/html",
            "investor_relations",
            "en",
        ),
        (pdf, "application/pdf", "investor_call_transcript", "unknown"),
        (
            "Full Conference Call Transcript\nCEO: 公司完成海外产能扩张，新产品已完成客户认证并进入量产。\n".encode(),
            "text/plain",
            "investor_call_transcript",
            "zh",
        ),
    ]
    with owned_lake() as (root, catalog):
        (root / "config").mkdir()
        cfg = root / "config/catalog.json"
        cfg.write_text(
            json.dumps(
                {
                    "schema_version": "1.0",
                    "catalog_dir": "catalog",
                    "roots": [
                        {
                            "root_id": "company_raw",
                            "path": "companies",
                            "kind": "company_raw",
                        }
                    ],
                }
            ),
            encoding="utf8",
        )
        results = [
            _public_official_import(
                root, cfg, data, mime=mime, kind=kind, language=language
            )
            for data, mime, kind, language in sources
        ]
        baseline = {
            p: p.read_bytes() for p in (root / "companies").rglob("*") if p.is_file()
        }
        template = {
            "schema_version": "narrative-batch-request/1",
            "run_id": "official-three",
            "profile": "P2",
            "max_seconds": 40,
            "max_tokens": 100000,
            "max_cost_usd": "1",
            "model": {
                "model_id": "stub-model",
                "endpoint": loopback_model_server.endpoint,
                "api_key_env": batch_fixtures.KEY_ENV,
                "max_output_tokens": 400,
                "timeout_seconds": 5,
                "allow_local_http": True,
            },
            "pricing": {
                "version": "fixture-price/1",
                "input_micro_usd_per_million_tokens": 1000000,
                "output_micro_usd_per_million_tokens": 2000000,
            },
        }
        batch = prepare_official_narrative_request(results, batch_template=template)
        req = root / "batch.json"
        req.write_text(json.dumps(batch.to_dict()), encoding="utf8")
        llm = root / "llm.yaml"
        llm.write_text(
            "llm:\n  provider: deepseek\n  model: stub-model\n  base_url: "
            + loopback_model_server.endpoint.removesuffix("/chat/completions")
            + "\n  max_tokens: 8192\n  temperature: 0.7\n",
            encoding="utf8",
        )
        call, output = _configured_invoke(root, cfg, req, llm)
        assert call.returncode == 0 and output["status"] == "completed", (
            call.stderr,
            output,
        )
        assert (
            len(loopback_model_server.requests) == 3
            and not loopback_model_server.errors
        )
        assert output["budget"] == {
            "tokens": 276,
            "estimated_micro_usd": 333,
            "unknown_reservations": 0,
            "unsettled_reservations": 0,
        }
        assert all(row["artifact_ref"] is not None for row in output["documents"])
        proofs = []
        for imported, (_data, _mime, kind, language) in zip(results, sources):
            source = imported["source_ref"]
            exact = SourceVersionReader(catalog).query_ref(
                source["document_id"], source["source_id"], source["content_sha256"]
            )
            assert (
                SourceVersionReader(catalog)
                .open_version(exact, purpose="source_export")
                .data
                == _data
            )
            ref, ref_receipt = _public_narrative(
                root,
                cfg,
                "reference",
                {
                    "schema_version": "narrative-reference-request/1",
                    "source_ref": imported["source_ref"],
                },
            )
            read, read_receipt = _public_narrative(
                root,
                cfg,
                "read",
                {
                    "schema_version": "narrative-read-request/1",
                    "narrative_ref": ref,
                    "as_of_date": "2026-10-09",
                    "expected_source": {
                        "canonical_entity_id": None,
                        "market": "US",
                        "security_id": "ACME",
                        "document_kind": kind,
                        "fiscal_year": None,
                        "fiscal_period": None,
                    },
                },
            )
            assert ref_receipt["status"] == "metadata_only"
            assert read_receipt["replay_status"] == "verified"
            assert read["summary"]["translate"] is False and read["summary"]["draft"][
                "language"
            ] == ("en" if language == "unknown" else language)
            assert read["evidence_spans"] and all(
                span["locator"] for span in read["evidence_spans"]
            )
            assert (
                read["source_ref"]["content_sha256"]
                == imported["source_ref"]["content_sha256"]
            )
            evidence_ids = {span["span_id"] for span in read["evidence_spans"]}
            assert all(
                set(claim["evidence_ids"]) <= evidence_ids
                for claim in read["summary"]["draft"]["claims"]
            )
            proofs.append(
                {
                    "source_ref": source,
                    "narrative_ref": ref,
                    "read_receipt": read_receipt,
                    "parser_versions": read["versions"],
                    "summary_language": read["summary"]["draft"]["language"],
                    "translate": False,
                    "locators": [
                        {"span_id": span["span_id"], "locator": span["locator"]}
                        for span in read["evidence_spans"]
                    ],
                    "selection": read["selection"],
                }
            )
        repeated_imports = [
            _public_official_import(
                root, cfg, data, mime=mime, kind=kind, language=language
            )
            for data, mime, kind, language in sources
        ]
        assert all(
            item["status"] == "deduplicated"
            and item["source_ref"] == original["source_ref"]
            for item, original in zip(repeated_imports, results)
        )
        again, resumed = _configured_invoke(root, cfg, req, llm)
        assert again.returncode == 0 and resumed["budget"] == output["budget"]
        assert len(loopback_model_server.requests) == 3
        assert (
            len(
                NarrativeRunStore(root / "AUTO.sqlite").reservations_for_run(
                    "official-three"
                )
            )
            == 3
        )
        assert baseline == {
            p: p.read_bytes() for p in (root / "companies").rglob("*") if p.is_file()
        }

        import os

        proof_file = os.environ.get("CWPA_W03_PROOF_FILE")
        if proof_file:
            Path(proof_file).write_text(
                json.dumps(
                    {
                        "schema_version": "w03-engineering-proof/1",
                        "fixture_only": True,
                        "supplier_requests": 0,
                        "supplier_cost_usd": "0",
                        "loopback_requests": len(loopback_model_server.requests),
                        "simulated_budget": output["budget"],
                        "repeat_budget": resumed["budget"],
                        "repeat_import_statuses": [
                            item["status"] for item in repeated_imports
                        ],
                        "canonical_bytes_and_sidecars_unchanged": True,
                        "documents": proofs,
                    },
                    ensure_ascii=False,
                    indent=2,
                )
                + "\n",
                encoding="utf-8",
            )


def test_image_only_pdf_is_parser_incomplete_not_successful_skip(
    loopback_model_server, monkeypatch
):
    import fitz
    from company_wiki.source_catalog.official_source_flow import (
        prepare_official_narrative_request,
    )

    monkeypatch.setenv("DEEPSEEK_API_KEY", batch_fixtures.KEY)
    doc = fitz.open()
    doc.new_page().draw_rect(
        fitz.Rect(70, 70, 200, 200), color=(0, 0, 0), fill=(0, 0, 0)
    )
    pdf = doc.tobytes()
    doc.close()
    with owned_lake() as (root, catalog):
        (root / "config").mkdir()
        cfg = root / "config/catalog.json"
        cfg.write_text(
            json.dumps(
                {
                    "schema_version": "1.0",
                    "catalog_dir": "catalog",
                    "roots": [
                        {
                            "root_id": "company_raw",
                            "path": "companies",
                            "kind": "company_raw",
                        }
                    ],
                }
            ),
            encoding="utf8",
        )
        imported = _import(catalog, pdf, mime="application/pdf", kind="prospectus")
        template = {
            "schema_version": "narrative-batch-request/1",
            "run_id": "image-official",
            "profile": "P2",
            "max_seconds": 15,
            "max_tokens": 10000,
            "max_cost_usd": "1",
            "model": {
                "model_id": "stub-model",
                "endpoint": loopback_model_server.endpoint,
                "api_key_env": batch_fixtures.KEY_ENV,
                "max_output_tokens": 400,
                "timeout_seconds": 5,
                "allow_local_http": True,
            },
            "pricing": {
                "version": "fixture-price/1",
                "input_micro_usd_per_million_tokens": 1000000,
                "output_micro_usd_per_million_tokens": 2000000,
            },
        }
        req = root / "batch.json"
        req.write_text(
            json.dumps(
                prepare_official_narrative_request(
                    [imported], batch_template=template
                ).to_dict()
            ),
            encoding="utf8",
        )
        llm = root / "llm.yaml"
        llm.write_text(
            "llm:\n  provider: deepseek\n  model: stub-model\n  base_url: "
            + loopback_model_server.endpoint.removesuffix("/chat/completions")
            + "\n  max_tokens: 8192\n  temperature: 0.7\n",
            encoding="utf8",
        )
        call, out = _configured_invoke(root, cfg, req, llm)
        assert out["status"] != "completed" and not loopback_model_server.requests
        assert "PARSER_INCOMPLETE" in json.dumps(out)
        assert out["documents"][0]["artifact_ref"] is None


def test_cn_fundraising_index_and_old_ipo_exclusion_are_honest():
    from company_wiki.source_catalog.official_discovery import (
        discover_official_documents,
    )

    body = '<html><article><time datetime="2026-08-15"></time><a href="/convertible.pdf">向不特定对象发行可转换公司债券募集说明书</a></article><article><time datetime="2021-08-15"></time><a href="/ipo.pdf">首次公开发行股票招股说明书</a></article></html>'.encode()
    out = discover_official_documents(
        [
            {
                "url": "https://official.example/cn-ir",
                "original": body,
                "mime_type": "text/html",
                "capture_id": "fixture-cn",
            }
        ],
        entity="另一家公司",
        as_of_date="2026-10-09",
        start_date="2026-02-01",
    )
    assert (
        len(out["candidates"]) == 1 and out["candidates"][0]["entity"] == "另一家公司"
    )
    assert out["candidates"][0]["document_kind"] == "convertible_bond_prospectus"
    assert (
        out["excluded"][0]["reason"] == "before_window"
        and out["no_documents_claimed"] is False
    )


def test_capture_offset_and_supplied_period_metadata_remain_exact():
    from company_wiki.source_catalog.official_source_flow import import_official_source

    body = b"<html><p>Company launched a new product and expanded capacity for customers.</p></html>"
    with owned_lake() as (root, catalog):
        value = request(body)
        value["capture_receipt"]["captured_at"] = "2026-10-09T02:30:00+02:00"
        value["source"].update(
            fiscal_year=2026,
            fiscal_period="Q2",
            period_end="2026-06-30",
            filing_date="2026-08-27",
        )
        result = import_official_source(catalog, original=body, request=value)
        assert result["metadata"]["retrieved_at"] == "2026-10-09T00:30:00Z"
        assert result["metadata"]["published_date"] == "2026-08-26"
        assert result["metadata"]["period_end"] == "2026-06-30"
        sidecar = json.loads(
            next((root / "companies").rglob("*.source.json")).read_text(
                encoding="utf-8"
            )
        )
        assert sidecar["filing_date"] == "2026-08-27"
        assert sidecar["fiscal_year"] == 2026 and sidecar["fiscal_period"] == "Q2"


@pytest.mark.parametrize("mode", ["truncated", "both"])
def test_discovery_cli_never_hides_omitted_local_indexes(mode):
    import os
    import subprocess
    import sys

    repo = Path(__file__).resolve().parents[2]
    body = b'<html><a data-date="2026-08-26" href="/one.pdf">Earnings call</a></html>'
    with owned_lake() as (root, _):
        first = root / "one.html"
        first.write_bytes(body)
        second = root / "two.html"
        second.write_bytes(body.replace(b"one", b"two"))
        value = {
            "schema_version": "official-discovery-request/1",
            "entity": "Acme",
            "start_date": "2026-02-01",
            "as_of_date": "2026-10-09",
            "max_bytes": len(body),
            "pages": [
                {
                    "input_file": str(path),
                    "url": "https://official.example/ir",
                    "mime_type": "text/html",
                    "capture_id": path.stem,
                }
                for path in (first, second)
            ],
        }
        if mode == "both":
            value["urls"] = []
        path = root / "discovery.json"
        path.write_text(json.dumps(value), encoding="utf-8")
        call = subprocess.run(
            [
                sys.executable,
                "-X",
                "utf8",
                "-B",
                "-m",
                "company_wiki.source_catalog.official_source_cli",
                "--operation",
                "discover",
                "--request",
                str(path),
            ],
            cwd=root,
            env={**os.environ, "PYTHONPATH": str(repo / "src")},
            capture_output=True,
            timeout=20,
        )
        if mode == "both":
            assert call.returncode == 2 and not call.stdout
        else:
            assert call.returncode == 0, call.stderr
            output = json.loads(call.stdout)
            assert output["coverage_status"] == "partial"
            assert "byte_limit_reached" in output["limitations"]
            assert output["no_documents_claimed"] is False


def test_reused_original_reports_the_persisted_publisher():
    from company_wiki.source_catalog.official_source_flow import import_official_source

    body = b"<html><p>Company launched a new product and expanded capacity for customers.</p></html>"
    with owned_lake() as (_, catalog):
        value = request(body)
        first = import_official_source(catalog, original=body, request=value)
        value["source"]["publisher"] = "Different capture publisher"
        again = import_official_source(catalog, original=body, request=value)
        assert again["source_ref"] == first["source_ref"]
        assert again["metadata"]["publisher"] == first["metadata"]["publisher"]


def test_long_cn_fundraising_business_and_routine_ir_use_existing_selection():
    from company_wiki.document_normalization import normalize_document
    from company_wiki.source_catalog.narrative_evidence import select_narrative_evidence
    from company_wiki.source_catalog.official_source_flow import import_official_source

    body = (
        "<html><h1>向不特定对象发行可转换公司债券募集说明书</h1>"
        + "".join(
            "<p>本文件第"
            + str(index)
            + "页介绍登记程序、资料递交、审批步骤和联系人信息。</p>"
            for index in range(120)
        )
        + "<h2>募集资金用途与海外业务</h2><p>公司完成海外产能扩张，新产品已完成客户认证并进入量产。</p></html>"
    ).encode()
    with owned_lake() as (_, catalog):
        value = request(
            body, entity="另一家公司", kind="convertible_bond_prospectus", language="zh"
        )
        value["source"].update(market="CN", security_id="688012")
        imported = import_official_source(catalog, original=body, request=value)
        ref = imported["source_ref"]
        document = normalize_document(
            body,
            source_id=ref["source_id"],
            source_sha256=ref["content_sha256"],
            mime_type="text/html",
        )
        selected = select_narrative_evidence(
            document.structure,
            title=value["source"]["title"],
            existing_kind="convertible_bond_prospectus",
        )
        assert (
            selected.status == "selected"
            and selected.document_kind == "convertible_bond_prospectus"
        )
        assert selected.coverage_complete and any(
            "海外产能扩张" in (span.raw_text or "") for span in selected.evidence_spans
        )
        assert selected.selected_text_bytes < len(body) // 5
        assert all(
            span.locator and span.source_id == ref["source_id"]
            for span in selected.evidence_spans
        )
        assert (
            imported["metadata"]["entity"] == "另一家公司"
            and imported["metadata"]["market"] == "CN"
        )
        routine = b"<html><p>Shareholder meeting registration contact and attendance procedure.</p></html>"
        routine_sha = hashlib.sha256(routine).hexdigest()
        from company_wiki.source_contract import source_id_for_sha256

        parsed = normalize_document(
            routine,
            source_id=source_id_for_sha256(routine_sha),
            source_sha256=routine_sha,
            mime_type="text/html",
        )
        skip = select_narrative_evidence(
            parsed.structure,
            title="Investor relations management policy",
            existing_kind="investor_relations",
        )
        assert (
            skip.status == "skipped_no_narrative"
            and skip.coverage_complete
            and not skip.evidence_spans
        )


def test_storage_failure_removes_only_this_calls_staged_original(monkeypatch):
    from company_wiki.source_catalog.official_source_flow import import_official_source
    from company_wiki.source_catalog.canonical_writer import (
        CanonicalSourceWriter,
        CanonicalImportError,
    )

    body = b"<html><p>Company launched a new product and expanded capacity for customers.</p></html>"
    with owned_lake() as (root, catalog):
        first = _import(catalog, body)
        baseline = {
            p: p.read_bytes() for p in (root / "companies").rglob("*") if p.is_file()
        }
        sentinel = root / "catalog/staging/other-owner.txt"
        sentinel.write_bytes(b"pre-existing staging sentinel")

        def fail(*_args, **_kwargs):
            raise CanonicalImportError("fixture_storage_failure")

        monkeypatch.setattr(CanonicalSourceWriter, "import_original_staged", fail)
        changed = body.replace(b"new product", b"new services")
        with pytest.raises(CanonicalImportError, match="fixture_storage_failure"):
            import_official_source(catalog, original=changed, request=request(changed))
        assert sentinel.read_bytes() == b"pre-existing staging sentinel"
        assert list(sentinel.parent.iterdir()) == [sentinel]
        assert baseline == {
            p: p.read_bytes() for p in (root / "companies").rglob("*") if p.is_file()
        }
        ref = first["source_ref"]
        exact = SourceVersionReader(catalog).query_ref(
            ref["document_id"], ref["source_id"], ref["content_sha256"]
        )
        assert (
            SourceVersionReader(catalog)
            .open_version(exact, purpose="source_export")
            .data
            == body
        )


@pytest.mark.parametrize("image_only", [False, True])
def test_official_presentation_import_preserves_original_and_opaque_content(image_only):
    import io
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    if image_only:
        from PIL import Image
        from pptx.util import Inches
        image = io.BytesIO()
        Image.new("RGB", (40, 30), "white").save(image, format="PNG")
        image.seek(0)
        slide.shapes.add_picture(image, Inches(0), Inches(0))
    else:
        from pptx.util import Inches
        slide.shapes.add_textbox(Inches(0), Inches(0), Inches(4), Inches(1)).text = (
            "Customer demand grows as the company expands production capacity.")
    original = io.BytesIO()
    deck.save(original)
    body = original.getvalue()
    mime = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    with owned_lake() as (root, catalog):
        first = _import(catalog, body, mime=mime, published=None)
        again = _import(catalog, body, mime=mime, published=None)
        assert first["source_ref"] == again["source_ref"] and again["status"] == "deduplicated"
        ref = first["source_ref"]
        reader = SourceVersionReader(catalog)
        exact = reader.query_ref(ref["document_id"], ref["source_id"], ref["content_sha256"])
        assert reader.open_version(exact, purpose="source_export").data == body
        assert len(list((root / "companies").rglob("*.pptx"))) == 1
        assert first["metadata"]["published_date"] is None
        # Successful original storage makes no claim about OCR/text coverage.
        assert "canonical_path" not in json.dumps(first) and str(root) not in json.dumps(first)


@pytest.mark.parametrize("body", [b"<html>not a presentation</html>", b"PK\x03\x04broken"])
def test_invalid_presentation_refused_before_canonical_write(body):
    from company_wiki.source_catalog.official_source_flow import OfficialSourceError
    mime = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    with owned_lake() as (root, catalog):
        with pytest.raises(OfficialSourceError, match="invalid_pptx"):
            _import(catalog, body, mime=mime)
        assert not list((root / "companies").rglob("*"))
