"""One OCR composition port: bounded snapshots, selected replay and causal identity."""

from dataclasses import replace
import hashlib
import importlib.util
import json
from pathlib import Path
from types import SimpleNamespace
import pytest
from company_wiki import document_normalization as dn
from company_wiki.automation.narrative_formats import parser_component
from company_wiki.automation.narrative_generation import generation_manifest
from company_wiki.automation.narrative_contracts import (
    NarrativeSelectResult,
    SourceRevisionEventPayload,
)
from company_wiki.automation.narrative_select import NarrativeSelectHandler
from company_wiki.automation.narrative_replay import replay_narrative_evidence
from company_wiki.automation.models import HandlerOutcome
from unit.test_narrative_select_handler import _payload, _context, FakeReader

MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def fixture_deck(**kwargs):
    path = Path(__file__).resolve().parents[1] / "document_normalization/conftest.py"
    spec = importlib.util.spec_from_file_location("owned_normalization_fixture", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.build_pptx(**kwargs)


def _module():
    from company_wiki.source_catalog import narrative_normalization

    return narrative_normalization


def config(root):
    root.mkdir(parents=True, exist_ok=True)
    models = {}
    for name in ("det", "cls", "rec"):
        p = root / (name + ".onnx")
        p.write_bytes(("owned tiny model " + name).encode())
        models[name] = {
            "path": str(p.resolve()),
            "sha256": hashlib.sha256(p.read_bytes()).hexdigest(),
        }
    return {
        "schema_version": "cwp-local-ocr-config/1",
        "models": models,
        "engine_version": "fixture-engine/1",
        "runtime_version": "fixture-runtime/1",
        "base_config_sha256": "a" * 64,
        "text_score": 0.5,
        "low_confidence_threshold": 0.8,
        "max_side_len": 3000,
        "intra_op_threads": 2,
        "inter_op_threads": 1,
    }


class FakeOCR:
    def __init__(self, config, *, confidence=0.95, blank=False):
        self.config = config
        self.identity_manifest = config.identity_manifest
        self.fingerprint = config.fingerprint
        self.calls = 0
        self.validations = 0
        self.confidence = confidence
        self.blank = blank

    def validate_environment(self):
        self.validations += 1
        for name in ("det", "cls", "rec"):
            getattr(self.config, name).verify()

    def transcribe(self, data, *, limits):
        limits.check_deadline()
        self.calls += 1
        return dn.ImageOCRResult(
            1,
            1,
            ()
            if self.blank
            else (
                dn.OCRLine(
                    "We launched a new product and expanded overseas production capacity for customers.",
                    (0.0, 0.0, 1.0, 1.0),
                    self.confidence,
                ),
            ),
        )


def port(root, **adapter_options):
    module = _module()
    instances = []

    def factory(value):
        adapter = FakeOCR(value, **adapter_options)
        instances.append(adapter)
        return adapter

    return module.NarrativeNormalization(
        dn.LocalOCRConfig.from_dict(config(root)), adapter_factory=factory
    ), instances


def normalize(value, data):
    return value.normalize(
        data,
        source_id="urn:company-wiki:source:sha256:" + hashlib.sha256(data).hexdigest(),
        source_sha256=hashlib.sha256(data).hexdigest(),
        mime_type=MIME,
    )


def test_actual_carrier_component_uses_pure_pptx_11_not_html_constant():
    assert parser_component("text/html")[1] == "1.0.0"
    assert parser_component(MIME)[1] == "1.1.0"


def test_generation_uses_actual_normalization_identity_and_overrides_parser_component():
    from unit.test_narrative_batch import _request, Reader, _module as batch

    request = _request()
    event = (
        batch()
        .build_batch_events(request, Reader(), now="2026-10-09T00:00:00Z")
        .events[0]
    )
    wire = json.loads(event.payload_json)
    wire["source_ref"]["mime_type"] = MIME
    wire["source_metadata"]["source_class"] = "filing"
    payload = SourceRevisionEventPayload.from_dict(wire)
    identity = {
        "parser_name": dn.PARSER_NAME,
        "parser_version": "2.0.0",
        "format": "pptx",
        "ocr_fingerprint": "b" * 64,
    }
    result = generation_manifest(
        request,
        payload,
        execution_versions=batch()._execution_versions(request),
        parser_components=identity,
    )
    assert result["parser_component"] == {"name": dn.PARSER_NAME, "version": "2.0.0"}
    assert result["parser_components"] == identity


def test_snapshot_is_frozen_pathless_causal_and_each_port_owns_its_adapter(tmp_path):
    module = _module()
    value, adapters = port(tmp_path / "models")
    identity = value.identity(MIME)
    snapshot = value.snapshot()
    assert adapters == [] and identity["parser_version"] == "2.0.0"
    assert str(tmp_path) not in json.dumps(identity)
    root = tmp_path / "project"
    (root / "config").mkdir(parents=True)
    (root / "config/local_ocr.json").write_text(json.dumps(snapshot), encoding="utf-8")
    frozen = module.NarrativeNormalization.from_project(
        root, enabled=True, adapter_factory=FakeOCR
    )
    (root / "config/local_ocr.json").write_text("{}", encoding="utf-8")
    assert frozen.snapshot() == snapshot
    child = module.NarrativeNormalization.from_snapshot(
        snapshot, adapter_factory=FakeOCR
    )
    assert frozen.identity(MIME) == child.identity(MIME)
    assert frozen.ocr_adapter is not child.ocr_adapter
    relocated = dn.LocalOCRConfig.from_dict(config(tmp_path / "moved"))
    assert relocated.fingerprint == value.config.fingerprint


def test_absent_config_is_pure_and_other_formats_ignore_broken_ocr_config(tmp_path):
    module = _module()
    assert (
        module.NarrativeNormalization.from_project(tmp_path, enabled=True).identity(
            MIME
        )["parser_version"]
        == "1.1.0"
    )
    (tmp_path / "config").mkdir()
    (tmp_path / "config/local_ocr.json").write_text("invalid", encoding="utf-8")
    pure = module.NarrativeNormalization.from_project(tmp_path, enabled=False)
    assert pure.identity("text/html") == dn.normalization_identity("text/html")


def test_ocr_selected_partial_is_usable_and_replays_one_unique_media(tmp_path):
    value, adapters = port(tmp_path / "models")
    data = fixture_deck(with_picture_only_slide=True, with_duplicate_picture=True)
    wire = _payload(
        data,
        title="Acme business update",
        document_kind="investor_relations",
        language="en",
        mime_type=MIME,
    )
    result = NarrativeSelectHandler(reader=FakeReader(wire, data), normalization=value)(
        _context(wire, lambda: None)
    )
    assert result.outcome is HandlerOutcome.SUCCEEDED, result.error
    selected = NarrativeSelectResult.from_dict(result.result)
    assert selected.parser.version == "2.0.0" and selected.selection.status == "partial"
    assert not selected.selection.coverage_complete and selected.evidence_spans
    assert any("ocr_used" in span.quality_flags for span in selected.evidence_spans)
    assert all(
        "low_ocr_confidence" not in span.quality_flags
        for span in selected.evidence_spans
    )
    assert adapters[0].calls == 1
    before = adapters[0].calls
    assert replay_narrative_evidence(data, selected, normalization=value) == len(
        selected.evidence_spans
    )
    assert adapters[0].calls - before == 1
    assert all(
        span.structured_value["language"] == "en" for span in selected.evidence_spans
    )


@pytest.mark.parametrize("version", ["1.0.0", "1.1.0"])
def test_historical_pure_replay_ignores_current_ocr_and_preserves_en(tmp_path, version):
    value, adapters = port(tmp_path / "models")
    data = fixture_deck()
    digest = hashlib.sha256(data).hexdigest()
    document = dn.normalize_document(
        data,
        source_id="urn:company-wiki:source:sha256:" + digest,
        source_sha256=digest,
        mime_type=MIME,
        parser_version=version,
    )
    span = replace(document.units[0], language="en").to_evidence_span(
        topics=["business_progress"], selection_reasons=["source_text"]
    )
    assert (
        value.replay(
            data,
            source_id="urn:company-wiki:source:sha256:" + digest,
            source_sha256=digest,
            mime_type=MIME,
            evidence_spans=[span],
        )
        == 1
    )
    assert adapters == []


def test_sha_mismatch_fails_before_any_image_inference(tmp_path):
    value, adapters = port(tmp_path / "models")
    Path(value.config.det.path).write_bytes(b"tampered owned model")
    with pytest.raises(dn.LocalOCRError, match="OCR_MODEL_SHA_MISMATCH"):
        normalize(value, fixture_deck(with_picture_only_slide=True))
    assert adapters[0].calls == 0


def image_deck(images=4):
    from io import BytesIO
    from pptx import Presentation
    from PIL import Image

    deck = Presentation()
    for index in range(images):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        pixels = BytesIO()
        Image.new("RGB", (1, 1), (index * 30, 0, 0)).save(pixels, format="PNG")
        slide.shapes.add_picture(BytesIO(pixels.getvalue()), 0, 0)
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


def test_original_language_samples_only_bounded_unique_images(tmp_path):
    from company_wiki.source_catalog.narrative_language import (
        detect_narrative_text_language,
    )

    value, adapters = port(tmp_path / "models")
    text = value.sample_text(image_deck(), MIME, max_images=2)
    assert detect_narrative_text_language(text) == "en"
    assert adapters[0].calls == 2 and adapters[0].validations == 1


@pytest.mark.parametrize("mode", ["blank", "low"])
def test_no_reliable_ocr_span_stays_incomplete_without_successful_skip(tmp_path, mode):
    value, adapters = port(
        tmp_path / "models",
        blank=mode == "blank",
        confidence=0.6 if mode == "low" else 0.95,
    )
    data = image_deck(1)
    wire = _payload(
        data,
        title="Company business update",
        document_kind="investor_relations",
        language="en",
        mime_type=MIME,
    )
    result = NarrativeSelectHandler(reader=FakeReader(wire, data), normalization=value)(
        _context(wire, lambda: None)
    )
    assert (
        result.outcome is HandlerOutcome.TERMINAL_FAILURE
        and result.error.code == "PARSER_INCOMPLETE"
    )
    assert not result.effects and result.metrics.tokens == 0
    assert adapters[0].calls == 1


def test_changed_config_fingerprint_cannot_replay_old_selected_media(tmp_path):
    value, adapters = port(tmp_path / "models")
    data = image_deck(1)
    document = normalize(value, data)
    spans = [
        replace(unit, language="en").to_evidence_span(
            topics=["business_progress"], selection_reasons=["original_text"]
        )
        for unit in document.units
    ]
    changed = _module().NarrativeNormalization(
        replace(value.config, text_score=0.7), adapter_factory=FakeOCR
    )
    with pytest.raises(dn.ReplayError, match="config identity"):
        changed.replay(
            data,
            source_id=document.source_id,
            source_sha256=document.source_sha256,
            mime_type=MIME,
            evidence_spans=spans,
            language="en",
        )
    assert changed.ocr_adapter.calls == 0


@pytest.mark.parametrize(
    "field,value",
    [
        ("low_confidence_threshold", 2),
        ("engine_version", ""),
        ("base_config_sha256", "bad"),
    ],
)
def test_invalid_explicit_config_is_refused_without_inference(tmp_path, field, value):
    wire = config(tmp_path / "models")
    wire[field] = value
    with pytest.raises(ValueError):
        _module().NarrativeNormalization.from_snapshot(wire)


@pytest.mark.parametrize("forged_generation", [False, True])
def test_partial_completed_derivation_public_replays_and_reuses(
    tmp_path, forged_generation
):
    from dataclasses import asdict
    from unit import test_narrative_verify_handler as verifying
    from company_wiki.automation.narrative_verify import NarrativeVerifyHandler
    from company_wiki.automation.narrative_contracts import NarrativeBundle
    from company_wiki.automation.narrative_generation import read_reuse_pin
    from company_wiki.source_catalog.source_reader import SourceRef
    from company_wiki.automation.models import canonical_json

    value, adapters = port(tmp_path / "models")
    data = image_deck(1)
    wire = _payload(
        data,
        title="Company business update",
        document_kind="investor_relations",
        language="en",
        mime_type=MIME,
    )
    selected_raw = NarrativeSelectHandler(
        reader=FakeReader(wire, data), normalization=value
    )(_context(wire, lambda: None))
    selected = NarrativeSelectResult.from_dict(selected_raw.result)
    summary = verifying._summary(selected)
    source_reader = verifying.FakeReader(selected, data)
    verified = NarrativeVerifyHandler(reader=source_reader, normalization=value)(
        verifying._context(selected, summary)
    )
    assert verified.outcome is HandlerOutcome.SUCCEEDED, verified.error
    bundle = NarrativeBundle.from_dict(verified.result)
    raw = canonical_json(bundle.to_dict()).encode()
    digest = hashlib.sha256(raw).hexdigest()
    ref = SourceRef(**bundle.source_ref.to_dict())
    manifest = {
        **asdict(ref),
        **source_reader.metadata,
        "canonical_entity_id": None,
        "market": "US",
        "security_id": "ACME",
        "fiscal_year": None,
        "fiscal_period": None,
    }
    source_reader.query_ref = lambda *args: ref
    source_reader.open_described_version = lambda current, **kwargs: (
        source_reader.opened,
        manifest,
    )
    version = SimpleNamespace(
        content_sha256=digest,
        byte_size=len(raw),
        selection_status="partial",
        quality_status=bundle.quality_status,
    )

    class Artifacts:
        def read_exact(self, **kwargs):
            assert kwargs["expected_sha256"] == digest and kwargs[
                "expected_size"
            ] == len(raw)
            return version, raw

    pin = {
        "artifact_version_id": "narrative-ocr-fixture",
        "content_sha256": digest,
        "byte_size": len(raw),
        "document_id": ref.document_id,
        "source_id": ref.source_id,
        "source_sha256": ref.content_sha256,
    }
    payload = SourceRevisionEventPayload.from_dict(wire)
    before = adapters[0].calls
    from unit.test_narrative_batch import _request
    from company_wiki.automation import narrative_batch as batch

    request = replace(
        _request(),
        sources=(payload.source_ref,),
        model_options_json=canonical_json(
            {**_request().model_options, "model_id": "fixture-v1"}
        ),
    )
    versions = {
        **batch._execution_versions(request),
        "adapter": "replay",
        "prompt": bundle.summary.model.prompt_version,
    }
    generation = generation_manifest(
        request,
        payload,
        execution_versions=versions,
        parser_components=value.identity(MIME),
    )
    if forged_generation:
        generation["parser_components"]["ocr_fingerprint"] = "f" * 64
        with pytest.raises(ValueError, match="BATCH_REUSE_BUNDLE_GENERATION_MISMATCH"):
            read_reuse_pin(
                Artifacts(),
                source_reader,
                payload,
                pin,
                manifest,
                generation,
                normalization=value,
            )
    else:
        assert read_reuse_pin(
            Artifacts(),
            source_reader,
            payload,
            pin,
            manifest,
            generation,
            normalization=value,
        )
    assert adapters[0].calls - before == 1
    assert (
        bundle.selection.status == "partial" and not bundle.selection.coverage_complete
    )
    assert bundle.summary.status == "completed" and bundle.summary.translate is False
    assert (
        bundle.evidence_spans[0].structured_value["normalization_quality"][
            "coverage_complete"
        ]
        is False
    )


# The production factory test's fixture supplies a real AUTO/catalog scope.
from unit.test_narrative_worker_factory import run_state  # noqa: E402,F401


def test_worker_factory_uses_frozen_snapshot_without_reading_current_config(
    run_state, tmp_path, monkeypatch  # noqa: F811 - imported pytest fixture
):
    from unit import test_narrative_worker_factory as factory_fixture

    factory = factory_fixture._factory()
    module = _module()
    captured = []
    wire = config(tmp_path / "models")
    options = {**run_state.options, "normalization_config": wire}
    local = run_state.root / "config/local_ocr.json"
    local.parent.mkdir(exist_ok=True)
    local.write_text("not the snapshot", encoding="utf-8")
    monkeypatch.setattr(
        module.NarrativeNormalization,
        "from_project",
        classmethod(lambda cls, *a, **k: pytest.fail("child reread deployment config")),
    )
    monkeypatch.setattr(dn, "LocalOCRAdapter", FakeOCR)
    register = factory.register_narrative_handlers

    def capture(registrar, dependencies):
        captured.append(dependencies.normalization)
        return register(registrar, dependencies)

    monkeypatch.setattr(factory, "register_narrative_handlers", capture)
    for _ in range(2):
        factory.create_runtime(factory_fixture._spec(run_state, options=options))
    assert captured[0].snapshot() == captured[1].snapshot() == wire
    assert (
        captured[0] is not captured[1]
        and captured[0].ocr_adapter is not captured[1].ocr_adapter
    )


from unit.test_narrative_versioned_resume import frozen_run  # noqa: E402,F401


def test_declared_binding3_without_manifest_is_rejected_not_downgraded(frozen_run):  # noqa: F811 - imported pytest fixture
    from company_wiki.automation import narrative_batch as batch
    from company_wiki.automation.models import canonical_json, canonical_json_hash
    from company_wiki.source_catalog.source_read_policy import (
        EXACT_READ_POLICY_FINGERPRINT_SCHEMA_VERSION,
    )

    state, request, store, runs, run = frozen_run
    frozen = json.loads(run.binding_json)
    policies = {
        ref.document_id: state.reader.read_policy_sha256(
            state.reader.query_ref(ref.document_id, ref.source_id, ref.content_sha256)
        )
        for ref in request.sources
    }
    frozen.update(
        schema_version="narrative-run-binding/3",
        read_policy_schema_version=EXACT_READ_POLICY_FINGERPRINT_SCHEMA_VERSION,
        source_read_policies=policies,
        read_policy_sha256=canonical_json_hash(policies),
        generation_settings={},
        generation_manifests={},
        reused_artifact_pins={},
    )
    broken = replace(run, binding_json=canonical_json(frozen))
    with pytest.raises(batch.BatchResumeError, match="BATCH_FROZEN_BINDING_INVALID"):
        batch._resume_binding(request, state.reader, store, runs, broken, deadline=None)
    assert runs.get_run(run.run_id) == run
    assert all(store.list_attempts(job_id) == () for job_id in run.job_ids)


def test_legacy_normalization_version_is_frozen_for_worker_replay():
    from company_wiki.automation import narrative_batch as batch
    from company_wiki.automation.models import canonical_json
    from unit.test_narrative_batch import _request, Reader

    binding = batch.build_batch_events(_request(), Reader(), now="2026-10-09T00:00:00Z")
    event = binding.events[0]
    wire = json.loads(event.payload_json)
    wire["source_ref"]["mime_type"] = MIME
    wire["source_metadata"]["source_class"] = "filing"
    binding = replace(
        binding, events=(replace(event, payload_json=canonical_json(wire)),)
    )
    assert batch._normalization_parsers(
        binding, {"document_normalization": "1.0.0"}
    ) == {event.subject_id: "1.0.0"}


@pytest.mark.parametrize("corruption", ["config", "manifest"])
def test_frozen_ocr_generation_rejects_config_identity_mismatch(tmp_path, corruption):
    from unit.test_narrative_batch import _request
    from company_wiki.automation import narrative_batch as batch

    value, adapters = port(tmp_path / "models")
    data = image_deck(1)
    wire = _payload(
        data,
        title="Business update",
        document_kind="investor_relations",
        language="en",
        mime_type=MIME,
    )
    payload = SourceRevisionEventPayload.from_dict(wire)
    request = replace(_request(), sources=(payload.source_ref,))
    generation = generation_manifest(
        request,
        payload,
        execution_versions=batch._execution_versions(request),
        parser_components=value.identity(MIME),
    )
    binding = batch._events_from_sources(
        request, [payload], (), now="2026-10-09T00:00:00Z"
    )
    snapshot = value.snapshot()
    binding = replace(
        binding,
        generation_manifests={payload.source_ref.document_id: generation},
        normalization_config=snapshot,
    )
    assert (
        json.loads(batch._frozen_binding(request, binding))["schema_version"]
        == "narrative-run-binding/3"
    )
    if corruption == "config":
        snapshot["text_score"] = 0.7
    else:
        generation["parser_components"]["ocr_fingerprint"] = "f" * 64
    with pytest.raises(batch.BatchResumeError, match="BATCH_FROZEN_BINDING_INVALID"):
        batch._frozen_binding(request, binding)
    assert adapters == []
