"""Tool-owned fixture config and retained native evidence; no OCR/provider calls."""
from contextlib import closing
import importlib
import json
import sqlite3
from types import SimpleNamespace

import pytest

from tools.cross_market_suite.runner import Suite

OCR = {"schema_version": "cwp-local-ocr-config/1", "engine_version": "3.8.1",
       "runtime_version": "1.26.0", "base_config_sha256": "a" * 64, "text_score": .5,
       "low_confidence_threshold": .8, "max_side_len": 3000, "intra_op_threads": 2,
       "inter_op_threads": 1, "models": {name: {"path": "/fixture/not-loaded/" + name,
       "sha256": "b" * 64} for name in ("det", "cls", "rec")}}


def configured_suite(tmp_path, mode, monkeypatch):
    from tools.cross_market_suite import runner
    spec, points = tmp_path / "cases.json", tmp_path / "checkpoints.json"
    spec.write_text(json.dumps({"cases": [{"case": "US-MSFT"}]}), encoding="utf-8")
    points.write_text(json.dumps({"per_company": ["format_pptx", "canonical_worker", "format_support", "narrative_virtual_read"],
                                  "shared": []}), encoding="utf-8")
    monkeypatch.setattr(runner, "SPEC", spec)
    monkeypatch.setattr(runner, "POINTS", points)
    host = tmp_path / "host"
    identities = host / "benchmarks/cross_market_rf"
    identities.mkdir(parents=True)
    for market in ("cn", "hk", "us"):
        (identities / ("identity_" + market + ".json")).write_text("{}", encoding="utf-8")
    monkeypatch.setattr(runner, "REPO", host)
    acquisition = host / "config/source_acquisition.yaml"
    acquisition.parent.mkdir(parents=True)
    acquisition.write_text(json.dumps({"staging_root": "${PROJECT_ROOT}/.source_catalog/staging",
                                       "adapters": {}}), encoding="utf-8")
    host_ocr = acquisition.with_name("local_ocr.json")
    host_ocr.write_text(json.dumps(OCR), encoding="utf-8")
    before = host_ocr.read_bytes()
    root = tmp_path / "owned"
    root.mkdir()
    suite = Suite(SimpleNamespace(mode=mode, suite="full", rf_root=tmp_path / "rf",
        ff_root=tmp_path / "ff", et_root=None, acquisition_config=acquisition,
        source_catalog_config=acquisition.with_name("source_catalog.yaml")), root)
    exports = []

    def export(source, name, parts):
        exports.append((name, list(parts)))
        target = root / name
        target.mkdir()
        if "config" in parts:
            (target / "config").mkdir()
            (target / "config/local_ocr.json").write_bytes(before)
        return target, "f" * 40

    monkeypatch.setattr(suite, "export", export)
    suite.setup()
    assert host_ocr.read_bytes() == before
    return suite, exports, before


def test_replay_declares_pure_parser_and_never_exports_host_ocr(tmp_path, monkeypatch):
    from company_wiki.source_catalog.narrative_normalization import NarrativeNormalization, PPTX_MIME
    suite, exports, _ = configured_suite(tmp_path, "replay", monkeypatch)
    assert "config" not in dict(exports)["wiki"]
    assert not (suite.wiki / "config/local_ocr.json").exists()
    port = NarrativeNormalization.from_project(suite.wiki, enabled=True)
    assert port.identity(PPTX_MIME)["parser_version"] == "1.1.0"
    assert port._adapter is None
    assert suite.context["normalization_policy"]["profile"] == "offline_fixture_no_ocr"


def test_live_keeps_explicit_current_head_config_snapshot(tmp_path, monkeypatch):
    from company_wiki.source_catalog.narrative_normalization import NarrativeNormalization, PPTX_MIME
    suite, exports, before = configured_suite(tmp_path, "live", monkeypatch)
    assert "config" in dict(exports)["wiki"]
    assert (suite.wiki / "config/local_ocr.json").read_bytes() == before
    port = NarrativeNormalization.from_project(suite.wiki, enabled=True)
    assert port.identity(PPTX_MIME)["parser_version"] == "2.0.0"
    assert port._adapter is None
    assert suite.context["normalization_policy"]["profile"] == "configured_head_snapshot"


def capture():
    try:
        module = importlib.import_module("tools.cross_market_suite.runtime_evidence")
    except ModuleNotFoundError:
        pytest.fail("missing tool-owned tiny native evidence capture")
    return module.capture_runtime_evidence


def native_db(tmp_path):
    # Native migrations; inserted rows are explicit synthetic unit fixtures, never provider receipts.
    from company_wiki.automation.store import AutomationStore
    db = AutomationStore(tmp_path / "auto.db").db_path
    with closing(sqlite3.connect(db)) as connection, connection:
        connection.execute("INSERT INTO narrative_runs "
            "(run_id,input_hash,scope_sha256,model_id,prompt_version,pricing_version,"
            "input_micro_usd_per_million_tokens,output_micro_usd_per_million_tokens,"
            "max_tokens,max_micro_usd,max_output_bytes,state,created_at,updated_at) "
            "VALUES ('fixture',?,?, 'stub','p1','r1',1,2,100,100,1000,'open','t0','t0')",
            ("a" * 64, "b" * 64))
        connection.execute("INSERT INTO jobs VALUES "
            "('j','jk','source.narrative_select','source_revision','s',?,'p1','1','low',"
            "'running',0,'t0',3,'event','t0','t0',NULL,NULL)", ("a" * 64,))
        connection.execute("INSERT INTO narrative_run_jobs VALUES ('fixture','j',?,'1')", ("a" * 64,))
        result = {"result": {"schema_version": "narrative-select-result/1",
                  "selection": {"status": "partial", "selected_count": 2, "coverage_complete": False},
                  "evidence_spans": [{"raw_text": "PRIVATE ORIGINAL TEXT"}]}}
        connection.execute("INSERT INTO attempts "
            "(attempt_id,job_id,attempt_no,worker_id,lease_token,lease_until,started_at,heartbeat_at,"
            "finished_at,outcome,result_json,error_code,error_detail,runtime_generation) VALUES "
            "('attempt','j',1,'private-worker','private-lease','t2','t0','t1',NULL,NULL,?,NULL,?,1)",
            (json.dumps(result), "PRIVATE DETAIL"))
        connection.execute("INSERT INTO narrative_model_reservations VALUES "
            "('attempt','fixture','j',?,20,30,100,40,'unknown',NULL,NULL,NULL,NULL,"
            "'MODEL_TIMEOUT',NULL,NULL,'t0',NULL,NULL)", ("c" * 64,))
        binding = {"schema_version": "narrative-run-binding/3", "normalization_config": None,
                   "generation_manifests": {"s": {"source_inputs": {
                       "parser_component": {"name": "document_normalization", "version": "1.1.0"}}}}}
        connection.execute("UPDATE narrative_runs SET binding_json=? WHERE run_id='fixture'", (json.dumps(binding),))
    return db


def test_tiny_evidence_retains_actual_partial_stage_unknown_and_artifact_null(tmp_path):
    db = native_db(tmp_path)
    receipt = {"status": "partial", "error": None, "documents": [{"status": "ready",
        "artifact_ref": None, "errors": []}], "budget": {"tokens": 50, "estimated_micro_usd": 40,
        "unknown_reservations": 1, "unsettled_reservations": 1}}
    before = db.read_bytes()
    evidence = capture()(db, receipt, run_id="fixture", returncode=2)
    assert db.read_bytes() == before
    assert evidence["receipt"]["status"] == "partial"
    assert evidence["receipt"]["documents"][0]["artifact_ref"] is None
    assert evidence["receipt"]["budget"]["unknown_reservations"] == 1
    assert evidence["normalization_binding"]["ocr_config_present"] is False
    assert evidence["normalization_binding"]["parsers"][0]["parser_version"] == "1.1.0"
    assert evidence["jobs"][0]["status"] == "running"
    assert evidence["jobs"][0]["latest_attempt"]["selection"]["selected_count"] == 2
    reservation = evidence["reservations"][0]
    assert reservation["usage_status"] == "unknown"
    assert reservation["input_tokens"] is None and reservation["charged_tokens"] == 50
    assert reservation["estimated_micro_usd"] is None and reservation["charged_micro_usd"] == 40
    encoded = json.dumps(evidence)
    assert len(encoded.encode()) <= 32768
    assert all(secret not in encoded for secret in ("PRIVATE", "private-worker", "private-lease", "raw_text"))
    db.unlink()
    assert evidence["jobs"][0]["latest_attempt"]["finished_at"] is None


def test_missing_database_is_unavailable_not_zero(tmp_path):
    evidence = capture()(tmp_path / "absent.db", {"status": "partial"}, run_id="absent")
    assert evidence["jobs"] is None and evidence["reservations"] is None
    assert evidence["receipt"]["budget"] is None
    assert evidence["database_evidence"] == "unavailable_database"
    assert not (tmp_path / "absent.db").exists()


def test_full_report_keeps_tiny_evidence_before_temp_cleanup(tmp_path, monkeypatch):
    suite, _, _ = configured_suite(tmp_path, "replay", monkeypatch)
    evidence = {"schema_version": "cmrf-runtime-evidence/1", "receipt": {"status": "partial"},
                "jobs": [{"job_type": "source.narrative_select", "status": "running"}],
                "reservations": [], "database_evidence": "available"}
    import xml.etree.ElementTree as ET
    test = ET.Element("testcase", name="format_case")
    props = ET.SubElement(test, "properties")
    for key, value in {"cmrf_case": "US-MSFT", "cmrf_format": "format_pptx",
                       "cmrf_format_detail": "actual partial",
                       "cmrf_runtime_evidence": json.dumps(evidence)}.items():
        ET.SubElement(props, "property", name=key, value=value)
    ET.SubElement(test, "failure")
    monkeypatch.setattr(suite, "command", lambda *a, **k: ET.ElementTree(test).write(suite.root / "worker.xml"))
    suite.full()
    check = suite.checks["US-MSFT", "format_pptx"]
    assert check["status"] == "FAIL"
    assert check["runtime_evidence"] == evidence



def test_single_image_pptx_public_cli_is_offline_blocked_with_actual_tiny_evidence(record_property):
    """One tiny fixture, actual CLI/worker, pure parser, zero HTTP/model reservations."""
    from datetime import datetime, timezone
    import hashlib
    import os
    from pathlib import Path
    import stat
    import subprocess
    import sys
    import time
    from PIL import Image, ImageDraw
    from pptx import Presentation
    from pptx.util import Inches
    from company_wiki.automation.models import RuntimeState
    from company_wiki.automation.store import AutomationStore
    from company_wiki.document_normalization import normalize_document
    from tools.cross_market_suite.core import format_capability_gap, isolated_directory, sha
    from tools.cross_market_suite.runtime_evidence import capture_runtime_evidence

    protected = Path(__file__).resolve().parents[2] / "config/local_ocr.json"
    protected_before = sha(protected) if protected.exists() else None
    with isolated_directory() as root:
        owned = root
        config = root / "config/source_catalog.yaml"
        config.parent.mkdir()
        config.write_text(json.dumps({"schema_version": "1.0", "catalog_dir": str(root / "catalog"),
            "roots": [{"root_id": "fixture_raw", "path": str(root / "companies"), "kind": "company_raw",
                       "adapter_id": "company_raw_v1", "read_only": False}]}), encoding="utf-8")
        assert not config.with_name("local_ocr.json").exists()
        raw = root / "input/image-only.pptx"
        raw.parent.mkdir(parents=True)
        image = root / "image.png"
        bitmap = Image.new("RGB", (240, 80), "white")
        ImageDraw.Draw(bitmap).text((10, 10), "Fixture product launch", fill="black")
        bitmap.save(image)
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(str(image), Inches(1), Inches(1))
        deck.save(raw)
        digest = hashlib.sha256(raw.read_bytes()).hexdigest()
        raw_before = (digest, raw.stat().st_size, raw.stat().st_mtime_ns)
        config_before = sha(config)
        raw.chmod(stat.S_IREAD)
        repo = Path(__file__).resolve().parents[2]
        env = {key: value for key, value in os.environ.items()
               if not any(word in key.upper() for word in ("KEY", "TOKEN", "SECRET", "PASSWORD"))}
        env.update(PYTHONPATH=os.pathsep.join((str(repo / "tools/cross_market_suite/replay_hook"), str(repo / "src"))),
                   PYTHONUTF8="1", PYTHONDONTWRITEBYTECODE="1", PYTHON_DOTENV_DISABLED="1", CMRF_OFFLINE="1",
                   OFFLINE_FIXTURE_UNUSED_KEY="synthetic-unused-fixture")
        import_request = {"schema_version": "official-source-import-request/1", "request_id": "tiny-fixture-import",
            "max_bytes": raw.stat().st_size, "content_sha256": digest,
            "mime_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "source": {"entity": "Fixture", "canonical_entity_id": "ent-offline-fixture", "market": "US",
                "security_id": "FIXTURE", "document_kind": "investor_relations", "language": "en",
                "fiscal_year": None, "fiscal_period": None, "title": "Synthetic fixture investor presentation",
                "publisher": "Fixture", "source_url": "https://fixtures.invalid/tiny-image-pptx"},
            "capture_receipt": {"capture_method": "local_document", "tool_name": "pytest_native_cli_fixture",
                "tool_call_id": "tiny-fixture-local-read", "captured_at": datetime.now(timezone.utc).isoformat(),
                "content_sha256": digest, "response_bytes": raw.stat().st_size}}
        import_path = root / "import-request.json"
        import_path.write_text(json.dumps(import_request), encoding="utf-8")
        import_argv = [sys.executable, "-B", "-m", "company_wiki.source_catalog.official_source_cli",
                       "--config", str(config), "--project-root", str(root), "--request", str(import_path),
                       "--input-file", str(raw)]
        imported = subprocess.run(import_argv, cwd=root, env=env, capture_output=True, timeout=20)
        assert imported.returncode == 0, imported.stderr
        registration = json.loads(imported.stdout)
        assert registration["download_events"] == 0
        ref = registration["source_ref"]
        record_property("cmrf_import_command", json.dumps(import_argv))
        record_property("cmrf_import_receipt", json.dumps({"status": registration["status"],
                        "source_ref": ref, "capture_receipt": registration["capture_receipt"], "download_events": 0}))
        work = root / "batch"
        work.mkdir()
        store = AutomationStore(work / "auto.db")
        store.set_runtime_gate(RuntimeState.PAUSED, updated_at="2026-10-09T00:00:00Z")
        request = {"schema_version": "narrative-batch-request/1", "run_id": "offline-tiny-pptx",
            "sources": [ref], "profile": "P2", "max_seconds": 30, "max_tokens": 1, "max_cost_usd": "0",
            "model": {"model_id": "fixture-unused-model", "endpoint": "http://127.0.0.1:9/v1/chat/completions",
                      "api_key_env": "OFFLINE_FIXTURE_UNUSED_KEY", "max_output_tokens": 1,
                      "timeout_seconds": 1, "allow_local_http": True},
            "pricing": {"version": "fixture-price/1", "input_micro_usd_per_million_tokens": 1,
                        "output_micro_usd_per_million_tokens": 1}}
        request_path = work / "request.json"
        request_path.write_text(json.dumps(request), encoding="utf-8")
        argv = [sys.executable, "-B", "-m", "company_wiki.automation.narrative_batch_cli",
                "--project-root", str(root), "--catalog-config", str(config), "--automation-db", str(store.db_path),
                "--work-dir", str(work / "jobs"), "--request", str(request_path)]
        start = time.monotonic()
        call = subprocess.run(argv, cwd=root, env=env, capture_output=True, timeout=40)
        result = json.loads(call.stdout)
        evidence = capture_runtime_evidence(store.db_path, result, run_id=request["run_id"], returncode=call.returncode)
        record_property("cmrf_runtime_evidence", json.dumps(evidence))
        record_property("cmrf_actual_command", json.dumps(argv))
        record_property("cmrf_seconds", str(round(time.monotonic() - start, 6)))
        structure = normalize_document(raw.read_bytes(), source_id=ref["source_id"], source_sha256=digest,
                                       mime_type=ref["mime_type"]).structure
        images_only = (not structure.units and not structure.errors and structure.page_count == structure.pages_read == 1
                       and structure.opaque_pages == (1,))
        # Existing strict classification still needs the actual opaque-page proof.
        assert images_only
        assert call.returncode == 2 and result["status"] == "failed", (call.stderr, result)
        assert format_capability_gap(result, images_only_checked=True) == "PARSER_INCOMPLETE"
        assert evidence["database_evidence"] == "available"
        assert evidence["normalization_binding"]["ocr_config_present"] is False
        assert {row["parser_version"] for row in evidence["normalization_binding"]["parsers"]} == {"1.1.0"}
        assert evidence["job_count"] == 3
        assert all(job["status"] == "dead_letter" for job in evidence["jobs"])
        assert evidence["reservation_count"] == 0 and evidence["reservations"] == []
        assert result["budget"] == {"tokens": 0, "estimated_micro_usd": 0,
                                    "unknown_reservations": 0, "unsettled_reservations": 0}
        assert all(doc["artifact_ref"] is None for doc in result["documents"])
        assert (sha(raw), raw.stat().st_size, raw.stat().st_mtime_ns) == raw_before
        assert sha(config) == config_before
        record_property("cmrf_raw_protection", json.dumps({"before": raw_before, "after": raw_before,
                         "config_before_sha256": config_before, "config_after_sha256": sha(config),
                         "read_only": not bool(raw.stat().st_mode & stat.S_IWRITE)}))
        record_property("cmrf_pure_structure", json.dumps({"pages_total": structure.page_count,
                        "pages_read": structure.pages_read, "source_units": len(structure.units),
                        "opaque_pages": list(structure.opaque_pages), "errors": list(structure.errors)}))
    assert not owned.exists()
    assert (sha(protected) if protected.exists() else None) == protected_before
    record_property("cmrf_protected_config", json.dumps({"before_sha256": protected_before, "after_sha256": protected_before}))
    record_property("cmrf_owned_root", str(owned))
    record_property("cmrf_owned_root_restored_absent", "true")



def test_legacy_binding_without_config_keeps_ocr_unknown(tmp_path):
    db = native_db(tmp_path)
    with closing(sqlite3.connect(db)) as connection, connection:
        connection.execute("UPDATE narrative_runs SET binding_json=?", (json.dumps({"schema_version": "narrative-run-binding/2"}),))
    evidence = capture()(db, {"status": "partial"}, run_id="fixture")
    assert evidence["normalization_binding"]["ocr_config_present"] is None
