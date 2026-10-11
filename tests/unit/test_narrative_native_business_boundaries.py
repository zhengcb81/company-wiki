"""Native structural boundaries are source order, never content-hash order."""
from dataclasses import replace
import hashlib
from io import BytesIO

import pytest

from company_wiki.document_normalization import normalize_document
from company_wiki.source_catalog.narrative_document import DocumentStructure
from company_wiki.source_catalog.narrative_evidence import (
    _make_unit, select_narrative_evidence,
)
from company_wiki.source_contract import EvidenceCoordinates, source_id_for_sha256
from support.native_document_fixtures import DOCX_MIME, docx_business_qa

BUSINESS = "Company launched a new product and expanded overseas production capacity."
CONTEXT = "The return calculation has not changed"
QUALIFIER = "Efficiency depends on workload mix, token usage and silicon price performance."
SHA = "a" * 64


def native_unit(text, kind, coordinates, **metadata):
    return replace(_make_unit(
        source_id=source_id_for_sha256(SHA), parser_version="1.0.0",
        coordinates=coordinates, raw_text=text, unit_kind=kind,
        source_role="company_filing", language="en", metadata=metadata,
    ), parser_name="cwp_document_normalization")


def select_units(units, version=None):
    structure = DocumentStructure(source_id=source_id_for_sha256(SHA),
        source_sha256=SHA, language="en", units=tuple(units), line_count=len(units))
    return select_narrative_evidence(structure, title="Company communication",
        existing_kind="investor_relations", selector_version=version)


def select_docx(data, version=None):
    sha = hashlib.sha256(data).hexdigest()
    doc = normalize_document(data, source_id=source_id_for_sha256(sha),
        source_sha256=sha, mime_type=DOCX_MIME)
    structure = replace(doc.structure, language="en", units=tuple(
        replace(unit, language="en") for unit in doc.units))
    return select_narrative_evidence(structure, title="Company communication",
        existing_kind="investor_relations", selector_version=version)


@pytest.mark.parametrize("kind", ["docx_table_cell", "html_table_cell", "pptx_table_cell"])
def test_native_financial_cell_cannot_be_reintroduced_as_sentence_context(kind):
    table = native_unit("12345", kind,
        EvidenceCoordinates(table_index=0, row_index=0, column_index=0),
        table_class="financial", table_headers=["Revenue"])
    business = native_unit(BUSINESS, kind.replace("table_cell", "paragraph"),
        EvidenceCoordinates(paragraph_index=0))
    result = select_units((table, business))
    assert [span.raw_text for span in result.evidence_spans] == [BUSINESS]
    assert result.dropped_financial_count == 1


def test_docx_table_interrupts_paragraph_completion_in_physical_body_order():
    from docx import Document
    document = Document()
    document.add_paragraph(CONTEXT)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Revenue"
    document.add_paragraph(QUALIFIER)
    out = BytesIO()
    document.save(out)
    result = select_docx(out.getvalue())
    assert {span.raw_text for span in result.evidence_spans} == {QUALIFIER}
    assert not any(span.structured_value.get("selection_group_id")
        for span in result.evidence_spans)
    document = Document()
    document.add_paragraph(CONTEXT)
    document.add_paragraph(QUALIFIER)
    out = BytesIO()
    document.save(out)
    control = select_docx(out.getvalue())
    assert {span.raw_text for span in control.evidence_spans} == {CONTEXT, QUALIFIER}
    assert len({span.structured_value.get("selection_group_id")
        for span in control.evidence_spans}) == 1


def test_native_business_table_cell_remains_independently_selected():
    from company_wiki.source_catalog.narrative_normalization import NarrativeNormalization
    data = docx_business_qa()
    sha = hashlib.sha256(data).hexdigest()
    doc = normalize_document(data, source_id=source_id_for_sha256(sha),
        source_sha256=sha, mime_type=DOCX_MIME)
    structure = NarrativeNormalization.language_structure(doc, "zh")
    result = select_narrative_evidence(structure, title="公司投资者关系活动记录",
        existing_kind="investor_relations")
    spans = [s for s in result.evidence_spans
        if s.raw_text == "新产品获得客户重复订单，并扩大海外生产产能。"]
    assert len(spans) == 1
    assert spans[0].structured_value["unit_kind"] == "docx_table_cell"
    assert not spans[0].structured_value.get("selection_group_id")
    assert spans[0].structured_value["source_locator"]


def test_pptx_text_shapes_are_separate_context_owners():
    a = native_unit(CONTEXT, "pptx_paragraph",
        EvidenceCoordinates(page_number=1, paragraph_index=0), shape_path="1")
    b = native_unit(QUALIFIER, "pptx_paragraph",
        EvidenceCoordinates(page_number=1, paragraph_index=0), shape_path="2")
    result = select_units((a, b))
    assert [span.raw_text for span in result.evidence_spans] == [QUALIFIER]
    same_shape = replace(b, metadata={"shape_path": "1"},
        coordinates=EvidenceCoordinates(page_number=1, paragraph_index=1))
    control = select_units((a, same_shape))
    assert {span.raw_text for span in control.evidence_spans} == {CONTEXT, QUALIFIER}


def test_new_policy_identity_preserves_explicit_old_native_selection():
    table = native_unit("12345", "docx_table_cell",
        EvidenceCoordinates(table_index=0, row_index=0, column_index=0),
        table_class="financial", table_headers=["Revenue"])
    business = native_unit(BUSINESS, "docx_paragraph",
        EvidenceCoordinates(paragraph_index=0))
    old = select_units((table, business), "0.7.0")
    assert {span.raw_text for span in old.evidence_spans} == {"12345", BUSINESS}
    new = select_units((table, business), "0.7.1")
    assert [span.raw_text for span in new.evidence_spans] == [BUSINESS]
    assert select_units((table, business), "0.6.0").evidence_spans


@pytest.mark.parametrize("kind", ["docx_heading", "html_heading"])
def test_native_headings_never_complete_other_heading_sentences(kind):
    first = native_unit(CONTEXT, kind,
        EvidenceCoordinates(paragraph_index=0), heading_level=1)
    second = native_unit(QUALIFIER, kind,
        EvidenceCoordinates(paragraph_index=1), heading_level=1)
    result = select_units((first, second))
    assert [span.raw_text for span in result.evidence_spans] == [QUALIFIER]
    assert not any(span.structured_value.get("selection_group_id")
        for span in result.evidence_spans)
