"""Formats share the real selection/publication path and bounded batch replay."""
from __future__ import annotations

from dataclasses import replace
from io import BytesIO
import importlib.util
from pathlib import Path
import sys

import pytest

from company_wiki.automation.models import HandlerOutcome
from company_wiki.automation.narrative_contracts import NarrativeBundle, NarrativeSelectResult
from company_wiki.automation.narrative_replay import NarrativeReplayError, replay_narrative_evidence
from company_wiki.document_normalization import PARSER_NAME, PARSER_VERSION
from company_wiki import document_normalization as dn
from company_wiki.source_contract import EvidenceSpan


def _fixture_module(name):
    spec = importlib.util.spec_from_file_location("r6_fixture_" + name, Path(__file__).with_name(name + ".py"))
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


selecting = _fixture_module("test_narrative_select_handler")
verifying = _fixture_module("test_narrative_verify_handler")

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
BUSINESS = "Company launched a new product and expanded overseas capacity for new customers."


def _pptx_bytes(*, image_only=False):
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    if image_only:
        from PIL import Image
        pixels = BytesIO()
        Image.new("RGB", (4, 4)).save(pixels, format="PNG")
        slide.shapes.add_picture(BytesIO(pixels.getvalue()), 0, 0, Inches(5), Inches(3))
    else:
        slide.shapes.add_textbox(0, 0, Inches(7), Inches(2)).text = BUSINESS
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


def _selected(data, mime):
    payload = selecting._payload(data, title="Company annual report", document_kind="annual_report",
                                 language="en", mime_type=mime)
    raw, _ = selecting._run(payload, data)
    assert raw.outcome is HandlerOutcome.SUCCEEDED, raw.error
    return NarrativeSelectResult.from_dict(raw.result)


@pytest.mark.parametrize("mime", ["text/html", "application/xhtml+xml", PPTX_MIME])
def test_formats_use_select_summary_and_verified_publication(mime):
    data = _pptx_bytes() if mime == PPTX_MIME else f"<html><body><p>{BUSINESS}</p></body></html>".encode()
    selected = _selected(data, mime)
    assert selected.parser.name == PARSER_NAME
    expected_version = "1.1.0" if mime == PPTX_MIME else "1.0.0"
    assert selected.parser.version == expected_version
    assert selected.selection.coverage_complete
    assert all(span.structured_value["language"] == "en" for span in selected.evidence_spans)
    result = verifying._run(data, selected, verifying._summary(selected))
    assert result.outcome is HandlerOutcome.SUCCEEDED, result.error
    bundle = NarrativeBundle.from_dict(result.result)
    assert bundle.replay.locator_count == len(selected.evidence_spans)
    assert bundle.versions.parser == expected_version
    assert replay_narrative_evidence(data, bundle) == len(selected.evidence_spans)


def test_html_financial_cells_are_excluded_but_business_cells_survive():
    data = ("<html><body><table><tr><th>Income statement</th><th>2025</th></tr>"
            "<tr><td>Revenue</td><td>12345</td></tr>"
            f"<tr><td>Operations</td><td>{BUSINESS}</td></tr></table></body></html>").encode()
    selected = _selected(data, "text/html")
    assert selected.selection.dropped_financial_count >= 4
    assert any(BUSINESS in (s.raw_text or "") for s in selected.evidence_spans)
    assert all(s.raw_text != "12345" for s in selected.evidence_spans)


@pytest.mark.parametrize("data,mime", [(b"not a package", PPTX_MIME), (b"<html><body></body></html>", "text/html")])
def test_broken_or_empty_format_cannot_be_accepted_as_no_narrative(data, mime):
    payload = selecting._payload(data, title="Company annual report", document_kind="annual_report",
                                 language="en", mime_type=mime)
    raw, _ = selecting._run(payload, data)
    assert raw.outcome is HandlerOutcome.TERMINAL_FAILURE
    assert raw.error.code == "PARSER_INCOMPLETE"
    assert raw.metrics.tokens == 0


def test_image_only_deck_is_named_incomplete_without_model_or_skip():
    data = _pptx_bytes(image_only=True)
    payload = selecting._payload(data, title="Company annual report", document_kind="annual_report",
                                 language="en", mime_type=PPTX_MIME)
    raw, _ = selecting._run(payload, data)
    assert raw.outcome is HandlerOutcome.TERMINAL_FAILURE
    assert raw.error.code == "PARSER_INCOMPLETE"
    assert raw.metrics.tokens == 0 and raw.metrics.cost_usd == 0


def test_all_html_spans_replay_with_one_parse_and_reject_forged_locator(monkeypatch):
    data = (f"<html><body><p>{BUSINESS}</p><p>We signed a new supply agreement with "
            "customers and expanded our production line.</p></body></html>").encode()
    selected = _selected(data, "text/html")
    assert len(selected.evidence_spans) >= 2
    normalizer = dn.normalize_document
    calls = []

    def counted(*args, **kwargs):
        calls.append(1)
        return normalizer(*args, **kwargs)

    monkeypatch.setattr(dn, "normalize_document", counted)
    assert replay_narrative_evidence(data, selected) == len(selected.evidence_spans)
    assert len(calls) == 1
    span = selected.evidence_spans[0]
    forged = EvidenceSpan.create(
        source_id=span.source_id, coordinates=span.coordinates, raw_text=span.raw_text,
        structured_value={**span.structured_value, "source_locator": "cwp-html-dom/1:missing"},
        parser_name=span.parser_name, parser_version=span.parser_version,
        parse_status=span.parse_status, quality_flags=span.quality_flags,
    )
    with pytest.raises(NarrativeReplayError):
        replay_narrative_evidence(data, replace(selected, evidence_spans=(forged,)))
    with pytest.raises(NarrativeReplayError):
        replay_narrative_evidence(data, replace(selected, parser=replace(selected.parser, version="9.9.9")))


def test_normalization_revision_is_frozen_only_for_new_format_batch_requests(monkeypatch):
    from company_wiki.automation import narrative_batch_request as requests
    fixtures = _fixture_module("test_narrative_batch_request")
    ordinary = requests.NarrativeBatchRequest.from_dict(fixtures._request())
    raw = fixtures._request()
    raw["sources"][0]["mime_type"] = PPTX_MIME
    normalized = requests.NarrativeBatchRequest.from_dict(raw)
    assert normalized.execution_versions["document_normalization"] == PARSER_VERSION
    assert "document_normalization" not in ordinary.execution_versions
    normalized_hash, ordinary_hash = normalized.input_hash, ordinary.input_hash
    monkeypatch.setattr(requests, "NORMALIZATION_PARSER_VERSION", "9.9.9")
    assert normalized.input_hash != normalized_hash
    assert ordinary.input_hash == ordinary_hash


@pytest.mark.parametrize("mime", ["text/html", "application/xhtml+xml", PPTX_MIME])
def test_batch_language_detection_uses_format_body_without_translation(mime):
    from company_wiki.source_catalog.narrative_language import detect_narrative_language
    data = _pptx_bytes() if mime == PPTX_MIME else f"<html><body><p>{BUSINESS}</p></body></html>".encode()
    assert detect_narrative_language(data, mime) == "en"


def test_html_transcript_keeps_its_transcript_parser_and_byte_lineage():
    from company_wiki.source_catalog.narrative_evidence import NARRATIVE_PARSER_NAME, NARRATIVE_PARSER_VERSION
    data = ("<html><body><p>Full Conference Call Transcript</p>"
            f"<p>CEO: {BUSINESS}</p></body></html>").encode()
    payload = selecting._payload(data, title="ACME call", document_kind="earnings_call_transcript",
                                 language="en", mime_type="text/html")
    raw, _ = selecting._run(payload, data)
    assert raw.outcome is HandlerOutcome.SUCCEEDED, raw.error
    selected = NarrativeSelectResult.from_dict(raw.result)
    assert (selected.parser.name, selected.parser.version) == (NARRATIVE_PARSER_NAME, NARRATIVE_PARSER_VERSION)
    assert selected.transcript_lineage is not None
    assert selected.transcript_byte_bindings
    result = verifying._run(data, selected, verifying._summary(selected))
    assert result.outcome is HandlerOutcome.SUCCEEDED, result.error
