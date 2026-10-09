"""W05 critical native format reachability and exact public evidence replay."""

from pathlib import Path
import importlib.util
import sys
from dataclasses import replace
import pytest
from company_wiki.automation.models import HandlerOutcome
from company_wiki.automation.narrative_contracts import (
    NarrativeBundle,
    NarrativeSelectResult,
)
from company_wiki.automation.narrative_replay import (
    NarrativeReplayError,
    replay_narrative_evidence,
)
from support.native_document_fixtures import (
    DOCX_MIME,
    docx_business_qa,
    natural_html_call,
)


def fixture(name):
    spec = importlib.util.spec_from_file_location(
        "w05_" + name, Path(__file__).with_name(name + ".py")
    )
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


selecting = fixture("test_narrative_select_handler")
verifying = fixture("test_narrative_verify_handler")


@pytest.mark.parametrize(
    "mime,kind,language,builder",
    [
        (DOCX_MIME, "investor_relations", "zh", docx_business_qa),
        ("text/html", "earnings_call_transcript", "en", natural_html_call),
    ],
)
def test_native_documents_select_verify_public_replay_and_roles(
    mime, kind, language, builder
):
    data = builder()
    payload = selecting._payload(
        data,
        title="Company business communication",
        document_kind=kind,
        language=language,
        mime_type=mime,
    )
    raw, _ = selecting._run(payload, data)
    assert raw.outcome is HandlerOutcome.SUCCEEDED
    selected = NarrativeSelectResult.from_dict(raw.result)
    assert selected.evidence_spans and selected.selection.coverage_complete
    assert any(
        s.structured_value["source_role"] == "management"
        for s in selected.evidence_spans
    )
    assert any(
        s.structured_value["source_role"] == "analyst" for s in selected.evidence_spans
    )
    company = next(
        s
        for s in selected.evidence_spans
        if s.structured_value["source_role"] == "management"
    )
    summary_input = replace(
        selected,
        evidence_spans=(
            company,
            *(s for s in selected.evidence_spans if s.span_id != company.span_id),
        ),
    )
    verified = verifying._run(data, selected, verifying._summary(summary_input))
    assert verified.outcome is HandlerOutcome.SUCCEEDED
    bundle = NarrativeBundle.from_dict(verified.result)
    assert bundle.summary.translate is False
    assert replay_narrative_evidence(data, bundle) == len(selected.evidence_spans)
    if mime == DOCX_MIME:
        assert (
            selected.parser.name == "cwp_document_normalization"
            and selected.parser.version == "1.0.0"
        )
        assert all(
            str(s.structured_value["source_locator"]).startswith("cwp-docx-body/1|")
            for s in selected.evidence_spans
        )
    else:
        assert selected.parser.version == "0.3.0" and selected.transcript_byte_bindings
        assert all(
            "unrelated retail" not in s.raw_text for s in selected.evidence_spans
        )
    with pytest.raises(NarrativeReplayError):
        replay_narrative_evidence(data + b"changed", bundle)
    assert all(
        s.structured_value["language"] == language for s in selected.evidence_spans
    )


def test_docx_financial_table_cells_stay_out_of_business_selection():
    from docx import Document
    from io import BytesIO

    document = Document()
    t = document.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Income statement"
    t.cell(0, 1).text = "Revenue"
    t.cell(1, 0).text = "Profit"
    t.cell(1, 1).text = "12345"
    document.add_paragraph(
        "Company launched a new product and expanded overseas production capacity."
    )
    out = BytesIO()
    document.save(out)
    data = out.getvalue()
    payload = selecting._payload(
        data,
        title="Company communication",
        document_kind="investor_relations",
        language="en",
        mime_type=DOCX_MIME,
    )
    raw, _ = selecting._run(payload, data)
    assert raw.outcome is HandlerOutcome.SUCCEEDED
    selected = NarrativeSelectResult.from_dict(raw.result)
    assert selected.selection.dropped_financial_count >= 4
    assert all(s.raw_text != "12345" for s in selected.evidence_spans)
