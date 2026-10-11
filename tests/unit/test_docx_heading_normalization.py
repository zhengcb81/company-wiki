"""Real OOXML heading structure, current DOCX identity and frozen 1.0 replay."""
from __future__ import annotations

from collections.abc import Mapping
from dataclasses import asdict
import hashlib
from io import BytesIO
import json
import zipfile

import pytest

from company_wiki import document_normalization as dn
from company_wiki.document_normalization.units import (
    parser_version_for_format, require_parser_version, verify_unit_identity,
)
from company_wiki.source_catalog.narrative_normalization import NarrativeNormalization
from company_wiki.source_contract import source_id_for_sha256

MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CONTEXT = "The return calculation has not changed"
HEADING = "Efficiency assumptions"
QUALIFIER = "Efficiency depends on workload mix, token usage and silicon price performance."
LEGACY_FINGERPRINT = "a6580fca5dcd6c063e2447bf5cf489aaabd8f61c52856646db744cb68d0031e4"
STYLES = (
    '<w:styles xmlns:w="' + W + '">'
    '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style>'
    '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/>'
    '<w:basedOn w:val="Normal"/><w:pPr><w:outlineLvl w:val="0"/></w:pPr></w:style>'
    '<w:style w:type="paragraph" w:styleId="Custom"><w:name w:val="Customer Section"/>'
    '<w:basedOn w:val="Heading1"/></w:style></w:styles>'
)


def paragraph(text, properties=""):
    return f"<w:p>{'<w:pPr>' + properties + '</w:pPr>' if properties else ''}<w:r><w:t>{text}</w:t></w:r></w:p>"


def package(body, styles=STYLES):
    parts = {
        "[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/></Types>',
        "_rels/.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
        "word/document.xml": f'<w:document xmlns:w="{W}"><w:body>{body}<w:sectPr/></w:body></w:document>',
        "word/_rels/document.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>',
    }
    if styles is not None:
        parts["word/styles.xml"] = styles
    out = BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for name, value in parts.items():
            archive.writestr(zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0)), value)
    return out.getvalue()


def legacy_original():
    return package(
        paragraph(CONTEXT)
        + paragraph(HEADING, '<w:pStyle w:val="Heading1"/>')
        + paragraph(QUALIFIER)
        + paragraph("Q: Has customer validation finished?")
        + paragraph("A: Customer validation remains planned next year.")
        + '<w:tbl><w:tr><w:tc>'
        + paragraph("New products received repeat customer orders.", '<w:pStyle w:val="Heading1"/>')
        + '</w:tc></w:tr></w:tbl>'
    )


def normalize(data, **kwargs):
    digest = hashlib.sha256(data).hexdigest()
    return dn.normalize_document(data, source_id=source_id_for_sha256(digest),
        source_sha256=digest, mime_type=MIME, **kwargs)


def plain(value):
    if isinstance(value, Mapping):
        return {key: plain(item) for key, item in value.items()}
    if isinstance(value, (tuple, list)):
        return [plain(item) for item in value]
    return value


def document_record(document):
    return {
        "source_sha256": document.source_sha256,
        "parser_version": document.parser_version,
        "metadata": plain(document.metadata),
        "errors": list(document.structure.errors),
        "units": [{
            "unit_id": unit.unit_id, "source_id": unit.source_id,
            "parser_name": unit.parser_name, "parser_version": unit.parser_version,
            "coordinates": asdict(unit.coordinates), "raw_text": unit.raw_text,
            "unit_kind": unit.unit_kind, "source_role": unit.source_role,
            "language": unit.language, "quality_flags": list(unit.quality_flags),
            "metadata": plain(unit.metadata),
        } for unit in document.units],
    }


def fingerprint(record):
    return hashlib.sha256(json.dumps(record, ensure_ascii=False, sort_keys=True,
        separators=(",", ":")).encode()).hexdigest()


def test_builtin_heading_from_real_word_style_keeps_paragraph_locator():
    from docx import Document
    document = Document()
    document.add_paragraph(CONTEXT)
    document.add_paragraph(HEADING, style="Heading 1")
    document.add_paragraph(QUALIFIER)
    out = BytesIO()
    document.save(out)
    doc = normalize(out.getvalue())
    heading = doc.units[1]
    assert heading.unit_kind == "docx_heading"
    assert heading.metadata["heading_level"] == 1
    assert heading.metadata["outline_level"] == 0
    assert heading.metadata["paragraph_style_id"] == "Heading1"
    assert heading.metadata["paragraph_style_name"] == "heading 1"
    assert heading.metadata["heading_basis"] == "style_outline"
    assert heading.metadata["source_locator"] == "cwp-docx-body/1|p=1"
    assert heading.coordinates.paragraph_index == 1 and heading.raw_text == HEADING


@pytest.mark.parametrize("properties,kind,level,basis", [
    ('<w:pStyle w:val="Custom"/>', "docx_heading", 1, "style_outline"),
    ('<w:outlineLvl w:val="3"/>', "docx_heading", 4, "paragraph_outline"),
    ('<w:pStyle w:val="Custom"/><w:outlineLvl w:val="9"/>', "docx_paragraph", None, None),
    ('<w:pStyle w:val="Custom"/><w:outlineLvl w:val="2"/>', "docx_heading", 3, "paragraph_outline"),
])
def test_inherited_style_and_direct_outline_take_real_precedence(properties, kind, level, basis):
    doc = normalize(package(paragraph("Identical wording", properties)))
    unit = doc.units[0]
    assert unit.unit_kind == kind
    assert unit.metadata.get("heading_level") == level
    assert unit.metadata.get("heading_basis") == basis
    if '<w:outlineLvl w:val="9"/>' in properties:
        assert unit.metadata["outline_level"] == 9
        assert unit.metadata["paragraph_style_id"] == "Custom"
    assert dn.replay_unit(package(paragraph("Identical wording", properties)),
        source_sha256=doc.source_sha256, unit=unit, limits=dn.NormalizationLimits()) == unit.raw_text


def test_style_name_and_heading_wording_do_not_guess_structure():
    styles = f'<w:styles xmlns:w="{W}"><w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/></w:style></w:styles>'
    doc = normalize(package(paragraph(HEADING, '<w:pStyle w:val="Heading1"/>'), styles))
    assert doc.units[0].unit_kind == "docx_paragraph"
    assert "heading_level" not in doc.units[0].metadata


def test_default_paragraph_style_uses_its_real_outline():
    styles = f'<w:styles xmlns:w="{W}"><w:style w:type="paragraph" w:default="1" w:styleId="DefaultSection"><w:name w:val="Default Section"/><w:pPr><w:outlineLvl w:val="1"/></w:pPr></w:style></w:styles>'
    unit = normalize(package(paragraph("Neutral wording"), styles)).units[0]
    assert unit.unit_kind == "docx_heading" and unit.metadata["heading_level"] == 2
    assert unit.metadata["paragraph_style_id"] == "DefaultSection"


def test_style_cycle_is_bounded_and_marks_incomplete_structure():
    styles = f'<w:styles xmlns:w="{W}"><w:style w:type="paragraph" w:styleId="A"><w:basedOn w:val="B"/></w:style><w:style w:type="paragraph" w:styleId="B"><w:basedOn w:val="A"/></w:style></w:styles>'
    doc = normalize(package(paragraph("Neutral wording", '<w:pStyle w:val="A"/>'), styles))
    assert doc.units[0].unit_kind == "docx_paragraph"
    assert "docx_style_inheritance_cycle" in doc.structure.errors
    assert not doc.structure.coverage_complete


def test_missing_style_and_invalid_outline_are_explicit_diagnostics():
    missing = normalize(package(paragraph("Neutral", '<w:pStyle w:val="Unknown"/>'), None))
    assert "docx_paragraph_style_missing" in missing.structure.errors
    invalid = normalize(package(paragraph("Neutral", '<w:outlineLvl w:val="10"/>')))
    assert invalid.units[0].unit_kind == "docx_paragraph"
    assert "docx_invalid_outline_level" in invalid.structure.errors


def test_styles_part_uses_existing_xml_limits_and_deadline():
    data = package(paragraph("Neutral"), '<broken')
    with pytest.raises(dn.UnsupportedFormatError, match="DOCX XML"):
        normalize(data)
    # The frozen parser never read styles, including a broken unused styles part.
    assert normalize(data, parser_version="1.0.0").units[0].raw_text == "Neutral"
    with pytest.raises(dn.NormalizationLimitError, match="deadline"):
        normalize(legacy_original(), limits=dn.NormalizationLimits(deadline=0))


def test_style_heading_inside_table_retains_table_semantics_and_body_order():
    doc = normalize(legacy_original())
    assert [unit.metadata["source_locator"] for unit in doc.units] == [
        "cwp-docx-body/1|p=0", "cwp-docx-body/1|p=1", "cwp-docx-body/1|p=2",
        "cwp-docx-body/1|p=3", "cwp-docx-body/1|p=4", "cwp-docx-body/1|t=0|r=0|c=0|p=0",
    ]
    assert doc.units[-1].unit_kind == "docx_table_cell"
    assert [unit.source_role for unit in doc.units[3:5]] == ["analyst", "management"]
    assert doc.units[3].metadata["qa_group_id"] == doc.units[4].metadata["qa_group_id"]


def test_explicit_legacy_is_exact_field_metadata_identity_and_replay_freeze():
    data = legacy_original()
    old = normalize(data, parser_version="1.0.0")
    assert fingerprint(document_record(old)) == LEGACY_FINGERPRINT
    assert old.units[1].unit_kind == "docx_paragraph"
    assert set(old.units[1].metadata) == {
        "normalization_schema", "format", "source_sha256", "source_locator", "transform",
    }
    assert dn.replay_units(data, source_sha256=old.source_sha256, units=old.units,
        limits=dn.NormalizationLimits()) == tuple(unit.raw_text for unit in old.units)


def test_default_and_saved_old_pins_have_separate_generation_and_replay_identity():
    from company_wiki.automation.narrative_formats import parser_component
    data = legacy_original()
    sha = hashlib.sha256(data).hexdigest()
    source_id = source_id_for_sha256(sha)
    port = NarrativeNormalization()
    pinned = NarrativeNormalization(parser_versions={"saved-document": "1.0.0"})
    assert dn.DOCX_PARSER_VERSION == "1.1.0"
    assert parser_version_for_format("docx") == "1.1.0"
    assert dn.normalization_identity(MIME)["parser_version"] == "1.1.0"
    assert port.identity(MIME)["parser_version"] == "1.1.0"
    assert pinned.identity(MIME, document_id="saved-document")["parser_version"] == "1.0.0"
    assert parser_component(MIME) == (dn.PARSER_NAME, "1.1.0")
    current = port.normalize(data, source_id=source_id, source_sha256=sha, mime_type=MIME)
    old = pinned.normalize(data, source_id=source_id, source_sha256=sha, mime_type=MIME,
        document_id="saved-document")
    assert fingerprint(document_record(old)) == LEGACY_FINGERPRINT
    assert current.parser_version == "1.1.0" and old.parser_version == "1.0.0"
    assert set(unit.unit_id for unit in current.units).isdisjoint(unit.unit_id for unit in old.units)
    for document in (old, current):
        for unit in document.units:
            verify_unit_identity(unit, format_name="docx")
        assert port.replay(data, source_id=source_id, source_sha256=sha, mime_type=MIME,
            evidence_spans=tuple(unit.to_evidence_span(topics=[], selection_reasons=[])
                for unit in document.units)) == len(document.units)
    with pytest.raises(ValueError, match="source_sha256"):
        dn.normalize_document(
            data + b"changed", source_id=source_id, source_sha256=sha, mime_type=MIME)


@pytest.mark.parametrize("format_name,version", [("docx", "2.0.0"), ("html", "1.1.0"), ("html", "2.0.0")])
def test_parser_capabilities_remain_specific_to_their_format(format_name, version):
    with pytest.raises(ValueError, match="parser_version"):
        require_parser_version(format_name, version)
    mime = MIME if format_name == "docx" else "text/html"
    with pytest.raises(ValueError, match="parser"):
        NarrativeNormalization().identity(mime, parser_version=version)


def test_real_heading_stops_native_business_group_context():
    from company_wiki.source_catalog.narrative_evidence import select_narrative_evidence
    doc = normalize(package(paragraph(CONTEXT)
        + paragraph(HEADING, '<w:pStyle w:val="Heading1"/>') + paragraph(QUALIFIER)))
    structure = NarrativeNormalization.language_structure(doc, "en")
    selected = select_narrative_evidence(structure, title="Company communication",
        existing_kind="investor_relations", selector_version="0.7.1")
    assert {span.raw_text for span in selected.evidence_spans} == {QUALIFIER}
    control = normalize(package(paragraph(CONTEXT) + paragraph(QUALIFIER)))
    selected_control = select_narrative_evidence(
        NarrativeNormalization.language_structure(control, "en"),
        title="Company communication", existing_kind="investor_relations", selector_version="0.7.1")
    assert {span.raw_text for span in selected_control.evidence_spans} == {CONTEXT, QUALIFIER}
