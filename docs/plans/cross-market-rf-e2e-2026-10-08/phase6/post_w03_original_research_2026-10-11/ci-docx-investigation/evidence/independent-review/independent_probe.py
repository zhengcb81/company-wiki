from __future__ import annotations
from pathlib import Path
from io import BytesIO
import sys, json, hashlib, datetime, traceback

ROOT = Path(r"C:\Users\郑曾波\.codex\worktrees\m3-acceptance-20261010\company-wiki")
OUT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "src"))
from company_wiki.document_normalization import normalize_document
from company_wiki.source_catalog.narrative_normalization import NarrativeNormalization
from company_wiki.source_catalog.narrative_evidence import select_narrative_evidence
from company_wiki.source_contract import source_id_for_sha256

BUSINESS = "Company launched a new product and expanded overseas production capacity."
CONTEXT = "The return calculation has not changed"
QUALIFIER = "Efficiency depends on workload mix, token usage and silicon price performance."
HEADING = "Efficiency assumptions"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
SOURCE_FILES = ["src/company_wiki/source_catalog/narrative_business_groups.py", "src/company_wiki/source_catalog/narrative_evidence.py"]
def sha(data): return hashlib.sha256(data).hexdigest()
def code_hash(): return {p: sha((ROOT/p).read_bytes()) for p in SOURCE_FILES}
def row(unit):
    return {"text": unit.raw_text, "unit_id": unit.unit_id, "kind": unit.unit_kind,
            "locator": unit.metadata.get("source_locator"), "coordinates": unit.coordinates.to_dict(),
            "shape_path": unit.metadata.get("shape_path"), "table_class": unit.metadata.get("table_class")}
def spanrow(span):
    return {"text": span.raw_text, "span_id": span.span_id, "locator": span.locator,
            "source_locator": span.structured_value.get("source_locator"),
            "kind": span.structured_value.get("unit_kind"),
            "group": span.structured_value.get("selection_group_id"),
            "reasons": span.structured_value.get("selection_reasons"),
            "shape_path": span.structured_value.get("shape_path"),
            "table_class": span.structured_value.get("table_class")}
def make_docx():
    from docx import Document
    document = Document()
    document.add_paragraph(BUSINESS)
    document.add_paragraph(CONTEXT)
    document.add_heading(HEADING, level=1)
    document.add_paragraph(QUALIFIER)
    table = document.add_table(rows=2, cols=2)
    for cell, text in zip([table.cell(0,0),table.cell(0,1),table.cell(1,0),table.cell(1,1)],
                          ["Revenue", "Net income", "12345", "67890"]): cell.text=text
    document.add_paragraph(BUSINESS)
    table = document.add_table(rows=2, cols=2)
    for cell,text in zip([table.cell(0,0),table.cell(0,1),table.cell(1,0),table.cell(1,1)],
                         ["Business", "Milestone", "New product", BUSINESS]): cell.text=text
    document.add_paragraph(CONTEXT)
    document.add_paragraph(QUALIFIER)
    out=BytesIO(); document.save(out); return out.getvalue()
def make_html():
    return ("<html><body><main><p>"+BUSINESS+"</p><p>"+CONTEXT+"</p><h2>"+HEADING+"</h2><p>"+QUALIFIER+"</p>"
        "<table><tr><th>Revenue</th><th>Net income</th></tr><tr><td>12345</td><td>67890</td></tr></table>"
        "<p>"+BUSINESS+"</p><table><tr><th>Business</th><th>Milestone</th></tr><tr><td>New product</td><td>"+BUSINESS+"</td></tr></table>"
        "<p>"+CONTEXT+"</p><p>"+QUALIFIER+"</p></main></body></html>").encode()
def make_pptx():
    from pptx import Presentation
    from pptx.util import Inches
    document = Presentation(); slide = document.slides.add_slide(document.slide_layouts[6])
    def shape(texts):
        tf=slide.shapes.add_textbox(Inches(0), Inches(0), Inches(5), Inches(1)).text_frame
        for i, text in enumerate(texts):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=text
    shape([CONTEXT]); shape([QUALIFIER]); shape([BUSINESS])
    table=slide.shapes.add_table(2,2, Inches(0), Inches(2), Inches(5), Inches(1)).table
    for cell,text in zip([table.cell(0,0),table.cell(0,1),table.cell(1,0),table.cell(1,1)],
                         ["Revenue", "Net income", "12345", "67890"]): cell.text=text
    table=slide.shapes.add_table(2,2, Inches(0), Inches(3), Inches(5), Inches(1)).table
    for cell,text in zip([table.cell(0,0),table.cell(0,1),table.cell(1,0),table.cell(1,1)],
                         ["Business", "Milestone", "New product", BUSINESS]): cell.text=text
    shape([CONTEXT, QUALIFIER])
    slide = document.slides.add_slide(document.slide_layouts[6]); shape([BUSINESS])
    out=BytesIO(); document.save(out); return out.getvalue()

report={"observed_at": datetime.datetime.now(datetime.timezone.utc).isoformat(), "source_before":code_hash(),
        "source_root":str(ROOT), "external_provider_calls":0,"external_model_calls":0,"probes":[]}
for name,mime,make in [("docx_heading_duplicates_mixed_tables",DOCX,make_docx),
                       ("html_heading_duplicates_mixed_tables","text/html",make_html),
                       ("pptx_same_and_different_shape_mixed_tables",PPTX,make_pptx)]:
    probe={"id":name,"checks":[]}
    try:
        data=make(); suffix={DOCX:"docx",PPTX:"pptx","text/html":"html"}[mime]
        (OUT/(name+"."+suffix)).write_bytes(data)
        digest=sha(data); source=source_id_for_sha256(digest)
        doc=normalize_document(data,source_id=source,source_sha256=digest,mime_type=mime)
        structure=NarrativeNormalization.language_structure(doc,"en")
        selected={v:select_narrative_evidence(structure,title="Company communication",existing_kind="investor_relations",selector_version=v)
                  for v in ["0.6.0","0.7.0","0.7.1"]}
        package=selected["0.7.1"]; spans=package.evidence_spans
        def check(check_id,success,detail): probe["checks"].append({"id":check_id,"status":"PASS" if success else "FAIL","detail":detail})
        probe.update(input_sha256=digest,input_bytes=len(data),native_units=[row(u) for u in doc.units],
                     selections={v:{"status":p.status,"dropped_financial_count":p.dropped_financial_count,"spans":[spanrow(s) for s in p.evidence_spans]} for v,p in selected.items()})
        business=[s for s in spans if s.raw_text==BUSINESS]
        check("repeated_text_keeps_distinct_original_locators",len(business)==3 and len({s.locator for s in business})==3,[s.locator for s in business])
        check("business_table_cell_standalone",any(s.raw_text==BUSINESS and s.structured_value.get("unit_kind")==suffix+"_table_cell" and not s.structured_value.get("selection_group_id") for s in spans),[spanrow(s) for s in business])
        check("financial_table_excluded",not any(s.raw_text in {"Revenue","Net income","12345","67890"} for s in spans),[spanrow(s) for s in spans if s.raw_text in {"Revenue","Net income","12345","67890"}])
        contexts=[s for s in spans if s.raw_text==CONTEXT]
        if suffix in {"docx","html"}:
            check("native_heading_not_sentence_context",not any(s.raw_text==HEADING for s in spans),[spanrow(s) for s in spans if s.raw_text==HEADING])
            check("adjacent_paragraph_positive_control",len(contexts)==1 and bool(contexts[0].structured_value.get("selection_group_id")),[spanrow(s) for s in contexts])
        else:
            check("shape_boundary_and_same_shape_positive_control",len(contexts)==1 and contexts[0].structured_value.get("shape_path")=="5" and bool(contexts[0].structured_value.get("selection_group_id")),[spanrow(s) for s in contexts])
        replay=NarrativeNormalization().replay(data,source_id=source,source_sha256=digest,mime_type=mime,evidence_spans=spans,language="en",parser_version=doc.parser_version)
        check("every_selected_original_locator_replays",replay==len(spans),{"replayed":replay,"selected":len(spans)})
    except Exception:
        probe["exception"]=traceback.format_exc(); probe["checks"].append({"id":"probe_execution","status":"FAIL","detail":probe["exception"]})
    probe["status"]="PASS" if all(c["status"]=="PASS" for c in probe["checks"]) else "FAIL"
    report["probes"].append(probe)
report["source_after"]=code_hash();report["source_unchanged"]=report["source_before"]==report["source_after"]
report["finished_at"]=datetime.datetime.now(datetime.timezone.utc).isoformat()
(OUT/"probe-results.json").write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print(json.dumps({"probes":[{"id":p["id"],"status":p["status"],"checks":p["checks"]} for p in report["probes"]],"source_unchanged":report["source_unchanged"]},ensure_ascii=False,indent=2))