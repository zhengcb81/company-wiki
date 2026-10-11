from pathlib import Path
from io import BytesIO
from dataclasses import replace
import json, hashlib, sys, datetime, zipfile
from lxml import etree
ROOT=Path(r"C:\Users\郑曾波\.codex\worktrees\m3-acceptance-20261010\company-wiki")
OUT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT/'src'))
from docx import Document
from company_wiki.document_normalization import normalize_document
from company_wiki.source_catalog.narrative_normalization import NarrativeNormalization
from company_wiki.source_catalog.narrative_evidence import select_narrative_evidence
from company_wiki.source_contract import source_id_for_sha256
CONTEXT='The return calculation has not changed'
HEADING='Efficiency assumptions'
QUALIFIER='Efficiency depends on workload mix, token usage and silicon price performance.'
doc=Document(); doc.add_paragraph(CONTEXT); doc.add_heading(HEADING,level=1); doc.add_paragraph(QUALIFIER)
out=BytesIO(); doc.save(out); data=out.getvalue(); digest=hashlib.sha256(data).hexdigest()
(OUT/'docx_heading_minimal.docx').write_bytes(data)
source_id=source_id_for_sha256(digest)
normalized=normalize_document(data,source_id=source_id,source_sha256=digest,mime_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
structure=NarrativeNormalization.language_structure(normalized,'en')
selected=select_narrative_evidence(structure,title='Company communication',existing_kind='investor_relations',selector_version='0.7.1')
with zipfile.ZipFile(BytesIO(data)) as package:
    xml=package.read('word/document.xml')
    node=etree.fromstring(xml); ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    raw=[{'text':''.join(p.xpath('.//w:t/text()',namespaces=ns)),
          'style':p.xpath('./w:pPr/w:pStyle/@w:val',namespaces=ns),
          'outline_level':p.xpath('./w:pPr/w:outlineLvl/@w:val',namespaces=ns)} for p in node.xpath('./w:body/w:p',namespaces=ns)]
spans=[{'text':s.raw_text,'locator':s.locator,'metadata':dict(s.structured_value)} for s in selected.evidence_spans]
report={'observed_at':datetime.datetime.now(datetime.timezone.utc).isoformat(),
        'input_sha256':digest,'input_bytes':len(data),'original_xml_sha256':hashlib.sha256(xml).hexdigest(),
        'raw_paragraphs':raw,
        'normalized_units':[{'text':u.raw_text,'kind':u.unit_kind,'metadata':dict(u.metadata)} for u in normalized.units],
        'selected_spans':spans,
        'expectation':'A true Heading 1 is a native structural boundary and must not be reintroduced as generic sentence context.',
        'status':'PASS' if not any(s.raw_text==HEADING for s in selected.evidence_spans) else 'FAIL',
        'replay_count':NarrativeNormalization().replay(data,source_id=source_id,source_sha256=digest,
            mime_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            evidence_spans=selected.evidence_spans,language='en',parser_version=normalized.parser_version),
        'source_sha256':{p:hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in [
            'src/company_wiki/source_catalog/narrative_business_groups.py','src/company_wiki/source_catalog/narrative_evidence.py']}}
(OUT/'heading-minimal-results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))